"""Servicio «Solo conversión»: estados de cuenta → Excel con resumen por ítem (WEI agosto, 3 bancos)."""
from decimal import Decimal as D

import openpyxl
import pytest

import test_wei_rest as tw
from cvp_parser.conversion import convertir, excel_conversion, items

pytestmark = pytest.mark.skipif(not tw.VENTAS.exists(), reason="faltan los archivos reales de WEI")


@pytest.fixture(scope="module")
def conv():
    return convertir("WEI REST", "2026-08", [(f"Cuenta {b}", b, edo) for b, edo, _ in tw.ARCHIVOS])


def test_lee_y_cuadra(conv):
    assert len(conv.cuentas) == 3 and conv.todo_cuadra
    for c in conv.cuentas:
        e = c.extracto
        # El resumen por ítem suma exactamente lo mismo que el estado de cuenta.
        its = items(c.movs)
        assert sum(i.debitos for i in its) == c.cuadre.total_debitos == sum(m.debito for m in e.movimientos)
        assert sum(i.creditos for i in its) == c.cuadre.total_creditos
        assert sum(i.n for i in its) == len(e.movimientos)
        assert e.saldo_anterior - c.cuadre.total_debitos + c.cuadre.total_creditos == c.cuadre.nuevo_saldo_calculado


def test_conceptos(conv):
    bnc = next(c for c in conv.cuentas if c.banco == "BNC")
    conceptos = {(i.naturaleza, i.concepto) for i in items(bnc.movs)}
    assert ("Comisiones bancarias", "Comisión Pago Movil") in conceptos
    traslados = [x for c in conv.cuentas for x in c.movs if x.naturaleza.startswith("Traslados")]
    assert len(traslados) == 12 and all(x.contraparte for x in traslados)


def test_excel(conv, tmp_path):
    salida = excel_conversion(conv, tmp_path / "conv.xlsx")
    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames[:2] == ["Resumen", "Resumen por ítem"] and len(wb.sheetnames) == 5
    ws = wb["Resumen por ítem"]
    textos = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Todas las cuentas" in textos and sum(t.startswith("TOTAL Cuenta") for t in textos) == 3
    # Cada débito y crédito del Excel es exactamente el del estado de cuenta (al céntimo).
    bnc = next(c for c in conv.cuentas if c.banco == "BNC")
    filas = list(wb["Mov Cuenta BNC"].iter_rows(min_row=5, max_row=4 + len(bnc.movs), values_only=True))
    leido = [(D(repr(f[5])) if f[5] is not None else D("0"), D(repr(f[6])) if f[6] is not None else D("0")) for f in filas]
    assert leido == [(x.mov.debito, x.mov.credito) for x in bnc.movs]

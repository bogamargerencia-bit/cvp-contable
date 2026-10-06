"""CACAO COFFEE NG, C.A., agosto 2026: Banco Plaza (cuenta personal de Johana Gil) + Banplus (.xls),
sistema en formato Mayor Analítico y cierre de caja con configuración propia (clientes.py).

Los archivos reales están en tests/fixtures/cacao_2026_08/ y no van en git; si faltan se omite.
"""
import datetime as dt
from collections import Counter
from decimal import Decimal as D
from pathlib import Path

import pytest

from cvp_parser.conciliacion import EstadoPartida as E
from cvp_parser.conciliacion_caja import EstadoCaja as EC, _dia_abono, fecha_venta
from cvp_parser.cuadre import Estado, TipoDiferencia
from cvp_parser.modelo import Movimiento
from cvp_parser.proceso import procesar_cliente
from cvp_parser.sistema import _es_ref_bancaria


def test_dia_abono_fin_de_semana_al_lunes():
    assert _dia_abono(dt.date(2026, 8, 1)) == dt.date(2026, 8, 3)    # sábado → lunes
    assert _dia_abono(dt.date(2026, 8, 2)) == dt.date(2026, 8, 3)    # domingo → lunes
    assert _dia_abono(dt.date(2026, 8, 4)) == dt.date(2026, 8, 4)    # martes


def test_fecha_venta():
    m = Movimiento(dt.date(2026, 8, 3), "Liquidacion Ventas Maestro", credito=D("1.00"),
                   referencia="783015200030056310776764")
    assert fecha_venta("BANPLUS", m) == dt.date(2026, 7, 31)
    m = Movimiento(dt.date(2026, 8, 20), "POS MAE 88046439 001 0003 1908", credito=D("1.00"))
    assert fecha_venta("PLAZA", m) == dt.date(2026, 8, 19)
    assert fecha_venta("BNC", m) is None


def test_referencia_bancaria():
    assert _es_ref_bancaria("083416407602")
    assert _es_ref_bancaria(83416407602.0)
    assert not _es_ref_bancaria(475000.0)
    assert not _es_ref_bancaria("PAGO MOVIL")


FIX = Path(__file__).parent / "fixtures" / "cacao_2026_08"
ARCH = [("PLAZA", FIX / "edo_plaza.pdf", FIX / "sistema_plaza.xls"),
        ("BANPLUS", FIX / "edo_banplus.xls", FIX / "sistema_banplus.xls")]
pytestmark_fix = pytest.mark.skipif(not all(p.exists() for _, a, b in ARCH for p in (a, b))
                                    or not (FIX / "ventas.xlsx").exists(),
                                    reason="faltan los archivos reales de CACAO (no van en git)")


@pytest.fixture(scope="module")
def rep():
    if not all(p.exists() for _, a, b in ARCH for p in (a, b)):
        pytest.skip("faltan los archivos reales de CACAO")
    return procesar_cliente("CACAO COFFEE NG, C.A.", "2026-08", ARCH, cierre_caja=FIX / "ventas.xlsx")


@pytestmark_fix
def test_plaza_detalle_completo_pero_resumen_impreso_inconsistente(rep):
    c = rep.cuadres["PLAZA"]
    e = c.extracto
    assert len(e.movimientos) == 159 and e.errores_lectura == []
    assert (c.total_debitos, c.total_creditos) == (D("10579695.84"), D("10633391.99"))
    tb = e.totales_banco
    assert (tb.nro_debitos, tb.nro_creditos) == (53, 106)        # conteos impresos coinciden
    tipos = {d.tipo for d in c.diferencias}
    # El saldo anterior impreso (3,97) y el «Saldo» final impreso no son coherentes con el detalle.
    assert tipos == {TipoDiferencia.SALDO_LINEA, TipoDiferencia.NUEVO_SALDO}
    assert e.movimientos[-1].saldo == D("53696.15")


@pytestmark_fix
def test_banplus_xls_cuadra_linea_por_linea(rep):
    c = rep.cuadres["BANPLUS"]
    assert c.diferencias == [] and c.estado is Estado.REQUIERE_REVISION
    assert len(c.extracto.movimientos) == 2304
    assert c.extracto.saldo_anterior == D("564173.84")
    assert c.nuevo_saldo_calculado == D("464032.88")


@pytestmark_fix
def test_mayor_analitico(rep):
    lp, lb = rep.libros["PLAZA"], rep.libros["BANPLUS"]
    assert (len(lp.asientos), lp.saldo_inicial_usd, lp.saldo_final_usd) == (16, D("0.00"), D("67.54"))
    assert (len(lb.asientos), lb.saldo_inicial_usd, lb.saldo_final_usd) == (154, D("656.79"), D("583.70"))
    assert lp.diferencias_saldo == [] and lb.diferencias_saldo == []
    assert any(a.referencia_banco == "083416407602" for a in lb.asientos)


@pytestmark_fix
def test_totales_del_mes_y_referencia(rep):
    def estado(b, texto):
        r = rep.conciliaciones[b]
        j = next(j for j, a in enumerate(r.libro.asientos) if texto in a.descripcion)
        return r.estado_asiento[j]
    p = estado("PLAZA", "INGRESO POR PUNTO DE VENTAS")
    assert p.estado is E.CONCILIADO_TOTAL and len(p.movs) == 106
    p = estado("PLAZA", "PAGO DE NOMINA 2DA QUINCENA")
    assert p.estado is E.CONCILIADO_TOTAL and len(p.movs) == 7 and "2.00" in p.nota
    p = estado("BANPLUS", "COMISIONES BANCARIAS BANPLUS")
    assert p.estado is E.CONCILIADO_TOTAL and len(p.movs) == 1014
    assert estado("BANPLUS", "AGUA MINERAL").estado is E.CONCILIADO_REFERENCIA
    p = estado("PLAZA", "COMISIONES BANCARIAS PLAZA")
    assert p.estado is E.SOLO_LIBRO and "1.30" in p.nota


@pytestmark_fix
def test_caja_cacao(rep):
    c = rep.caja
    assert "La fecha 27/08/2026 aparece en 2 filas (30, 31)." in c.cierre.errores
    assert "No hay fila para el 26/08/2026." in c.cierre.errores
    (t,) = c.totales_mes
    assert (t.caja, t.abonado, t.mes_anterior, t.por_abonar) == (
        D("44928569.08"), D("43651664.55"), D("938962.12"), D("2215866.65"))
    pm = Counter(l.estado for l in c.lineas if l.medio == "Pago móvil")
    assert pm[EC.CONCILIADO] == 3
    fila30 = next(l for l in c.lineas if l.medio == "Pago móvil" and l.fecha == dt.date(2026, 8, 27))
    assert "fila 30" in fila30.nota and "26/08" in fila30.nota
    cruces = {(x.banco, x.medio): x for x in c.asientos}
    assert cruces[("BANPLUS", "Pago móvil")].conciliado
    pos = cruces[("BANPLUS", "Punto Cacao + Punto Johana Gil")]
    assert not pos.conciliado and pos.usd_libro_total == D("50803.10") and pos.usd_caja_mes == D("51685.12")


@pytestmark_fix
def test_correccion_de_fecha_de_caja_en_tres_corridas(tmp_path):
    import openpyxl
    from cvp_parser.exportar import excel_conciliacion
    from cvp_parser.revision import Situacion, TIPO_CORRECCION, aplicar, correcciones_caja, leer_revision

    def correr(salida, revision=None):
        ant = leer_revision(revision) if revision else None
        rep = procesar_cliente("CACAO COFFEE NG, C.A.", "2026-08", ARCH, cierre_caja=FIX / "ventas.xlsx",
                               correcciones_caja=correcciones_caja(ant) if ant else None)
        rv = aplicar(rep, ant)
        excel_conciliacion(rep, salida, rv)
        return rep, rv

    rep1, rv1 = correr(tmp_path / "c1.xlsx")
    (pr,) = rep1.caja.propuestas
    assert (pr.fila, pr.fecha_actual, pr.fecha_propuesta) == (30, dt.date(2026, 8, 27), dt.date(2026, 8, 26))
    assert any(p.tipo == TIPO_CORRECCION for p in rv1.pendientes)

    # El analista acepta la corrección.
    wb = openpyxl.load_workbook(tmp_path / "c1.xlsx")
    ws = wb["Revisión"]
    r0 = next(r for r in range(1, 40) if ws.cell(r, 1).value == "Código")
    h = {ws.cell(r0, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(r0 + 1, ws.max_row + 1):
        if ws.cell(r, h["Tipo"]).value == TIPO_CORRECCION:
            ws.cell(r, h["Decisión"], "Aceptar")
            ws.cell(r, h["Revisado por"], "Ana")
    wb.save(tmp_path / "c1d.xlsx")

    rep2, rv2 = correr(tmp_path / "c2.xlsx", tmp_path / "c1d.xlsx")
    c = rep2.caja
    assert c.cierre.correcciones == [(30, dt.date(2026, 8, 27), dt.date(2026, 8, 26))]
    assert c.cierre.errores == [] and c.propuestas == []
    dia26 = next(l for l in c.lineas if l.medio == "Pago móvil" and l.fecha == dt.date(2026, 8, 26))
    assert dia26.estado is EC.CONCILIADO
    assert not any("26/08 sin día" in m for _, _, m in c.sin_caja)
    resueltas = {d.tipo for d in rv2.resueltas}
    assert TIPO_CORRECCION in resueltas and "Cierre de caja con errores" in resueltas

    # Corrida 3 con el archivo de la corrida 2 (sin tocar nada): la corrección se sigue aplicando.
    rep3, rv3 = correr(tmp_path / "c3.xlsx", tmp_path / "c2.xlsx")
    assert rep3.caja.cierre.correcciones == [(30, dt.date(2026, 8, 27), dt.date(2026, 8, 26))]
    assert not any(p.tipo == TIPO_CORRECCION for p in rv3.pendientes)
    assert all(p.situacion is not Situacion.NUEVO for p in rv3.pendientes)

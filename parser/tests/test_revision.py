"""Ciclo de revisión: corrida 1 → decisiones del analista → corrida 2 → ... → OK general.

Usa los archivos reales de WEI REST (se omite si no están, porque no van en git).
"""
from pathlib import Path

import openpyxl
import pytest
import xlrd

from cvp_parser.exportar import excel_conciliacion
from cvp_parser.proceso import procesar_cliente
from cvp_parser.revision import DecisionLeida, Situacion, _validar, aplicar, leer_revision, pendientes

FIX = Path(__file__).parent / "fixtures" / "wei_rest_2026_08"
NECESARIOS = ["edo_100_banco.pdf", "sistema_100_banco.xls", "edo_bnc.xls", "sistema_bnc.xls",
              "edo_banplus.xlsx", "sistema_banplus.xls", "ventas.xlsx"]


def test_validar_decisiones():
    d = lambda dec, com="", rev="Ana": DecisionLeida("X", dec, com, rev, "", "")
    assert _validar(d("Aceptar")) is None
    assert _validar(d("Justificado", "en tránsito")) is None
    assert "comentario" in _validar(d("Justificado"))
    assert "Revisado por" in _validar(d("Aceptar", rev=""))
    assert "no válida" in _validar(d("Listo"))


needs_fix = pytest.mark.skipif(not all((FIX / n).exists() for n in NECESARIOS),
                               reason="faltan los archivos reales de WEI REST (no van en git)")


def correr(sistema_100: Path, salida: Path, revision: Path | None = None):
    rep = procesar_cliente("WEI REST", "2026-08", [
        ("100_BANCO", FIX / "edo_100_banco.pdf", sistema_100),
        ("BNC", FIX / "edo_bnc.xls", FIX / "sistema_bnc.xls"),
        ("BANPLUS", FIX / "edo_banplus.xlsx", FIX / "sistema_banplus.xls"),
    ], cierre_caja=FIX / "ventas.xlsx")
    rv = aplicar(rep, leer_revision(revision) if revision else None)
    excel_conciliacion(rep, salida, rv)
    return rep, rv


def decidir(archivo: Path, salida: Path, regla) -> None:
    """Simula al analista: regla(tipo, descripcion, situacion) -> (decision, comentario) o None."""
    wb = openpyxl.load_workbook(archivo)
    ws = wb["Revisión"]
    r0 = next(r for r in range(1, 40) if ws.cell(r, 1).value == "Código")
    h = {ws.cell(r0, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(r0 + 1, ws.max_row + 1):
        if not ws.cell(r, 1).value or ws.cell(r, h["Tipo"]).value is None:
            continue
        res = regla(ws.cell(r, h["Tipo"]).value, ws.cell(r, h["Descripción"]).value,
                    ws.cell(r, h["Situación"]).value)
        if res:
            ws.cell(r, h["Decisión"], res[0])
            ws.cell(r, h["Comentario"], res[1])
            ws.cell(r, h["Revisado por"], "Ana")
    wb.save(salida)


def corregir_sistema_100(destino: Path) -> Path:
    """Export del sistema de 100% con la referencia 62222.81 corregida a 62222.80, guardado como .xlsx."""
    sh = xlrd.open_workbook(str(FIX / "sistema_100_banco.xls")).sheet_by_index(0)
    wb = openpyxl.Workbook()
    for r in range(sh.nrows):
        v = sh.row_values(r)
        if r and sh.cell_type(r, 0) == xlrd.XL_CELL_DATE:
            v[0] = xlrd.xldate_as_datetime(v[0], 0)
        if v[2] == 62222.81:
            v[2] = 62222.80
        wb.active.append([None if x == "" else x for x in v])
    wb.save(destino)
    return destino


@needs_fix
def test_codigos_estables_entre_corridas():
    rep1, _ = correr(FIX / "sistema_100_banco.xls", Path("/dev/null"))
    rep2, _ = correr(FIX / "sistema_100_banco.xls", Path("/dev/null"))
    c1 = [p.codigo for p in pendientes(rep1)]
    assert c1 == [p.codigo for p in pendientes(rep2)]
    assert len(c1) == len(set(c1))


@needs_fix
def test_ciclo_completo(tmp_path):
    rep, rv1 = correr(FIX / "sistema_100_banco.xls", tmp_path / "c1.xlsx")
    assert rv1.corrida == 1 and not rv1.ok_general
    assert all(p.situacion is Situacion.NUEVO for p in rv1.pendientes)

    otro = {"hecho": False}

    def regla1(tipo, desc, _):
        if tipo == "Diferencia de monto" and desc == "ECOMAS, COMIDA PERSONAL":
            return ("Corregido en el sistema", "Se corrigió a 62.222,80")
        if tipo == "Diferencia de monto":
            return ("Aceptar", None)
        if tipo.startswith("Caja: En tránsito"):
            return ("Justificado", "El lote se abona en septiembre")
        if tipo == "Lote sin día de caja":
            return ("Justificado", None)                         # falta el comentario
        if tipo == "Registrado en el libro de otro banco" and not otro["hecho"]:
            otro["hecho"] = True
            return ("Corregido en el sistema", None)             # pero no se corrige
        return None

    decidir(tmp_path / "c1.xlsx", tmp_path / "c1d.xlsx", regla1)
    sist = corregir_sistema_100(tmp_path / "sistema_100.xlsx")
    _, rv2 = correr(sist, tmp_path / "c2.xlsx", tmp_path / "c1d.xlsx")
    from collections import Counter
    cnt = Counter(p.situacion for p in rv2.pendientes)
    assert rv2.corrida == 2
    assert [d.descripcion for d in rv2.resueltas] == ["ECOMAS, COMIDA PERSONAL"]
    assert cnt[Situacion.CERRADO] == 25
    assert cnt[Situacion.INCOMPLETA] == 2
    assert cnt[Situacion.CORREGIDO_SIGUE] == 1
    assert cnt[Situacion.NUEVO] == 0
    n_hist = len(rv2.historial)

    # Corrida 3 sin tocar nada: el historial no se duplica.
    _, rv3 = correr(sist, tmp_path / "c3.xlsx", tmp_path / "c2.xlsx")
    assert rv3.corrida == 3 and len(rv3.historial) == n_hist

    # Corrida 4: se justifica todo lo abierto → OK general.
    decidir(tmp_path / "c3.xlsx", tmp_path / "c3d.xlsx",
            lambda t, d, s: ("Justificado", "Revisado con el cliente") if s != "Cerrado" else None)
    _, rv4 = correr(sist, tmp_path / "c4.xlsx", tmp_path / "c3d.xlsx")
    assert rv4.ok_general and not rv4.abiertas
    ws = openpyxl.load_workbook(tmp_path / "c4.xlsx")["Revisión"]
    assert ws["A3"].value.startswith("OK GENERAL")


@needs_fix
def test_revision_de_otro_cliente_se_rechaza(tmp_path):
    _, _ = correr(FIX / "sistema_100_banco.xls", tmp_path / "c1.xlsx")
    rep = procesar_cliente("OTRO", "2026-08", [("BNC", FIX / "edo_bnc.xls", FIX / "sistema_bnc.xls")])
    with pytest.raises(ValueError, match="no de OTRO"):
        aplicar(rep, leer_revision(tmp_path / "c1.xlsx"))

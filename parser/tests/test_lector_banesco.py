"""Banesco, estado de cuenta en PDF, agosto 2026 (INVERSIONES P.T.S.)."""
import datetime as dt
from decimal import Decimal as D
from pathlib import Path

import pytest

from cvp_parser.conversion import convertir, items
from cvp_parser.cuadre import Estado, cuadrar_extracto
from cvp_parser.lectores import LECTORES

PDF = Path(__file__).parent / "fixtures" / "banesco_2026_08" / "edo_banesco.pdf"
SIN_PDF = pytest.mark.skipif(not PDF.exists(), reason="falta el estado de cuenta real de Banesco")


@SIN_PDF
def test_lee_y_cuadra():
    assert LECTORES["BANESCO"].reconoce(PDF)
    assert not [k for k, l in LECTORES.items() if k != "BANESCO" and l.reconoce(PDF)]
    [e] = LECTORES["BANESCO"].leer([PDF]).extractos
    assert (e.cuenta, e.desde, e.hasta, e.titular) == ("01340025330251067692", dt.date(2026, 8, 1),
                                                       dt.date(2026, 8, 31), "INVERSIONES P.T.S. C,A.")
    assert e.saldo_anterior == D("35330.60") and not e.errores_lectura
    assert len(e.movimientos) == 10 and all(m.saldo is not None for m in e.movimientos)
    q = cuadrar_extracto(e)
    assert q.estado is Estado.CUADRA, [d.mensaje for d in q.diferencias]
    assert (q.total_debitos, q.total_creditos, q.nuevo_saldo_calculado) == (
        D("35850.01"), D("285854.41"), D("285335.00"))
    m = e.movimientos[2]
    assert (m.fecha_contable, m.referencia, m.descripcion, m.credito) == (
        dt.date(2026, 8, 18), "03617656953", "TRANS,CTAS", D("10122.63"))


@SIN_PDF
def test_conversion():
    conv = convertir("PTS", "2026-08", [("Banesco 7692", "BANESCO", PDF)])
    assert conv.todo_cuadra
    its = {(i.naturaleza, i.concepto): i for i in items(conv.cuentas[0].movs)}
    assert its[("Cargos bancarios (mantenimiento, estado de cuenta, SMS)", "EMISION DE ESTADO DE CUENTA")].debitos == D("54.00")
    assert sum(i.creditos for i in its.values() if i.naturaleza == "Transferencias recibidas") == D("285854.41")



@pytest.mark.parametrize("desc", ["S.MAST.29082026-821908302", "S.DISG.31082026-821908303", "BCV.C.O.TPV.08/26-821908303"])
def test_comisiones_punto_de_venta_en_cualquier_banco(desc):
    """Inversiones PTS: comisiones por marca de tarjeta y cargo BCV por terminal = comisiones (todo banco)."""
    from cvp_parser.cargos import tipo_de
    from cvp_parser.modelo import Movimiento
    from cvp_parser.naturaleza import S_COMISION, clasificar
    m = Movimiento(dt.date(2026, 8, 31), desc, debito=D("794.99"))
    assert all(clasificar(b, m) == S_COMISION for b in ("BANESCO", "BNC", "OTRO_BANCO"))
    assert tipo_de(m) in ("S.MAST", "S.DISG", "BCV.C.O.TPV")


@SIN_PDF
def test_resumen_del_banco_en_el_excel(tmp_path):
    """La columna izquierda del estado (resumen de movimientos y de saldos) también llega al Excel."""
    import openpyxl
    from cvp_parser.conversion import excel_conversion
    [e] = LECTORES["BANESCO"].leer([PDF]).extractos
    assert e.totales_banco.lineas == [
        ("Cheques Pagados", 0, D("0.00")), ("Otros Débitos", 6, D("35850.01")), ("Depósitos", 0, D("0.00")),
        ("Otros Créditos", 4, D("285854.41")), ("Saldo Inicial", None, D("35330.60")),
        ("Saldo Final", None, D("285335.00")), ("Saldo Promedio", None, D("19799.86"))]
    out = excel_conversion(convertir("PTS", "2026-08", [("Banesco 7692", "BANESCO", PDF)]), tmp_path / "c.xlsx")
    ws = openpyxl.load_workbook(out)["Mov Banesco 7692"]
    col_d = [ws.cell(r, 4).value for r in range(1, ws.max_row + 1)]
    assert "Resumen del banco (tal como aparece en el estado de cuenta)" in col_d
    i = col_d.index("Otros Créditos") + 1
    assert (ws.cell(i, 5).value, D(repr(ws.cell(i, 6).value))) == (4, D("285854.41"))
    assert "Comparación: banco vs. detalle convertido" in col_d

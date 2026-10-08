"""Lector del Kardex de tesorería en divisas (WEI REST, «KARDEX AGOSTO 2026»).

Estructura observada (hoja del mes):
  - Fila de encabezado: FECHA | DIA | ZELLE | Bolivares | DIVISAS | USDT || FECHA | PAGOS DIVISAS | MONTO ||
    FECHA | PAGOS BOLIVARES | MONTO | TASA. Montos en US$ (la columna Bolivares también, convertida).
  - Primera fila «VIENEN»: saldo inicial por medio.
  - Filas por día: ingresos del día por medio. Una fila puede seguir sin fecha (pertenece al último día).
    Si en la fila hay un texto («COMPRA», «REPARTO», «CXC PAGADAS», «PINTURA»…), los montos de esa fila
    no son ventas sino un movimiento con esa descripción. Un monto negativo sin texto es una salida.
  - Bloques a la derecha: pagos en divisas (salen de DIVISAS) y pagos en bolívares (informativo).
  - Al pie: sección «FONDO DE EFECTIVO» (VIENEN, y filas Nro/Fecha/Monto) y «DISPONIBLE DIVISAS / ZELLE / USDT».

Todos los montos se convierten con monto_excel (2 decimales exactos o error).
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Optional

import openpyxl

from .montos import CERO, MontoInvalido, monto_excel

MEDIOS = ["ZELLE", "BOLIVARES", "DIVISAS", "USDT"]


@dataclass
class MovKardex:
    fila: int
    fecha: Optional[dt.date]
    medio: str                 # ZELLE / BOLIVARES / DIVISAS / USDT / FONDO
    monto: Decimal             # + entra, − sale
    descripcion: str
    tipo: str                  # venta_dia / evento / pago_divisas / pago_bolivares / fondo


@dataclass
class Kardex:
    archivo: str
    hoja: str
    movimientos: list[MovKardex]
    saldo_inicial: dict[str, Decimal]
    disponible: dict[str, Decimal]
    errores: list[str] = field(default_factory=list)

    def ventas_dia(self, medio: str) -> dict[dt.date, Decimal]:
        out: dict[dt.date, Decimal] = {}
        for m in self.movimientos:
            if m.tipo == "venta_dia" and m.medio == medio:
                out[m.fecha] = out.get(m.fecha, CERO) + m.monto
        return out

    def no_ventas(self, medio: str) -> list[MovKardex]:
        return [m for m in self.movimientos if m.medio == medio and m.tipo != "venta_dia"]


def _n(t: object) -> str:
    return re.sub(r"\s+", " ", str(t or "")).strip().upper()


def _fecha_texto(t: Optional[str], anio: int) -> Optional[dt.date]:
    """«07-09», «7/9», «07-09-2026» → fecha (día-mes; el año, si falta, es el del Kardex)."""
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})(?:[-/.](\d{2}|\d{4}))?", (t or "").strip())
    if not m:
        return None
    a = int(m.group(3)) if m.group(3) else anio
    a = a + 2000 if a < 100 else a
    try:
        return dt.date(a, int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def leer_kardex(ruta: Path, hoja: Optional[str] = None) -> Kardex:
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".pdf":           # Kardex impreso a PDF (Shiro): ver kardex_pdf.py
        from .kardex_pdf import leer_kardex_pdf
        return leer_kardex_pdf(ruta)
    wb = openpyxl.load_workbook(str(ruta), data_only=True)
    ws = wb[hoja] if hoja else next(w for w in wb.worksheets if w.sheet_state == "visible")
    errores: list[str] = []

    f_enc = next(r for r in range(1, 20) if {"ZELLE", "DIVISAS", "USDT"} <= {_n(c.value) for c in ws[r]})
    enc = {}
    for c in ws[f_enc]:
        t = _n(c.value)
        if t:
            enc.setdefault(t, []).append(c.column)
    col = {m: enc[m][0] for m in ("ZELLE", "BOLIVARES", "DIVISAS", "USDT")}
    c_fecha = enc["FECHA"][0]
    c_pd = enc["PAGOS DIVISAS"][0]
    c_pb = enc["PAGOS BOLIVARES"][0]

    def monto(v) -> Optional[Decimal]:
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            try:
                return monto_excel(v)
            except MontoInvalido as e:
                errores.append(str(e))
        return None

    movs: list[MovKardex] = []
    saldo_ini: dict[str, Decimal] = {}
    disponible: dict[str, Decimal] = {}
    fin_dias = next((r for r in range(f_enc + 1, ws.max_row + 1)
                     if _n(ws.cell(r, c_fecha).value) in ("INGRESOS MENSUALES", "TOTALES")), ws.max_row)
    actual: Optional[dt.date] = None
    for r in range(f_enc + 1, fin_dias):
        f = ws.cell(r, c_fecha).value
        textos = [_n(ws.cell(r, c).value) for c in range(c_fecha, col["USDT"] + 1)
                  if isinstance(ws.cell(r, c).value, str) and _n(ws.cell(r, c).value)]
        if "VIENEN" in textos:
            for m, c in col.items():
                v = monto(ws.cell(r, c).value)
                if v is not None:
                    saldo_ini[m] = v
        else:
            if isinstance(f, dt.datetime):
                actual = f.date()
            etiqueta = " ".join(t for t in textos if t not in DIAS)
            for m, c in col.items():
                v = monto(ws.cell(r, c).value)
                if v is None or v == 0 or actual is None:
                    continue
                if etiqueta:
                    movs.append(MovKardex(r, actual, m, v, etiqueta.title(), "evento"))
                elif v < 0:
                    movs.append(MovKardex(r, actual, m, v, "Salida sin descripción", "evento"))
                else:
                    movs.append(MovKardex(r, actual, m, v, "Ventas del día", "venta_dia"))
        # Bloques de pagos (misma fila).
        for c_desc, medio, tipo in ((c_pd, "DIVISAS", "pago_divisas"), (c_pb, "BOLIVARES", "pago_bolivares")):
            desc, fv, v = ws.cell(r, c_desc).value, ws.cell(r, c_desc - 1).value, monto(ws.cell(r, c_desc + 1).value)
            if isinstance(desc, str) and desc.strip() and v is not None and isinstance(fv, dt.datetime):
                movs.append(MovKardex(r, fv.date(), medio, -v, desc.strip(), tipo))

    # Pie: DISPONIBLE ... y sección FONDO DE EFECTIVO.
    for r in range(fin_dias, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            t = _n(ws.cell(r, c).value)
            m = re.match(r"DISPONIBLE (DIVISAS|ZELLE|USDT|BOLIVARES)$", t)
            if m:
                v = monto(ws.cell(r, c + 1).value)
                if v is not None:
                    disponible[m.group(1)] = v
    f_fondo = next((r for r in range(fin_dias, ws.max_row + 1)
                    if any("FONDO DE EFECTIVO" in _n(c.value) for c in ws[r])), None)
    if f_fondo:
        anio = actual.year if actual else dt.date.today().year
        for r in range(f_fondo + 1, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, 8)]
            textos = [v.strip() for v in vals if isinstance(v, str) and v.strip()]
            t = [_n(v) for v in textos]
            fecha = next((v.date() for v in vals if isinstance(v, dt.datetime)), None)
            fecha_txt = None
            if fecha is None:            # fecha escrita como texto («07-09», «07/09/2026»), en cualquier columna
                fecha_txt = next((x for x in textos if _fecha_texto(x, anio)), None)
                fecha = _fecha_texto(fecha_txt, anio) if fecha_txt else None
            nums = [monto(v) for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            nums = [n for n in nums if n is not None]
            if "VIENEN" in t and nums:
                saldo_ini["FONDO"] = nums[-1]
            elif fecha and nums:
                etiqueta = next((_n(x) for x in textos if x != fecha_txt and _n(x) != "NRO."), "")
                movs.append(MovKardex(r, fecha, "FONDO", nums[-1], f"Fondo de efectivo {etiqueta}".strip(), "fondo"))
                if fecha_txt:
                    errores.append(f"Fondo de efectivo, fila {r}: la fecha está escrita como texto («{fecha_txt}»); "
                                   f"se leyó como {fecha:%d/%m/%Y}.")
            elif not fecha and nums and not t and "FONDO" in saldo_ini and "FONDO" not in disponible:
                disponible["FONDO"] = nums[-1]          # total al pie de la sección
                break                                   # lo que sigue son notas, no movimientos
        # Control: VIENEN + movimientos = total al pie. Si no, hay una fila que no se pudo leer.
        if "FONDO" in saldo_ini and "FONDO" in disponible:
            suma = sum((m.monto for m in movs if m.medio == "FONDO"), CERO)
            calc = saldo_ini["FONDO"] + suma
            if calc != disponible["FONDO"]:
                errores.append(f"La sección FONDO DE EFECTIVO no cuadra: VIENEN {saldo_ini['FONDO']} + movimientos "
                               f"{suma} = {calc}, pero el total al pie es {disponible['FONDO']} (diferencia "
                               f"{disponible['FONDO'] - calc}). Hay alguna fila que no se pudo leer (¿sin fecha?).")
    return Kardex(ruta.name, ws.title, movs, saldo_ini, disponible, errores)


DIAS = {"LUNES", "LINES", "MARTES", "MIERCOLES", "MIÉRCOLES", "JUEVES", "VIERNES", "SABADO", "SÁBADO", "DOMINGO"}

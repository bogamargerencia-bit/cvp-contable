"""Banplus — export .xlsx de movimientos ("Cuenta-<número>-Des...").

Formato observado (agosto 2026, cuenta 0174-0124661244397772):
  - Columnas: Fecha, Referencia (con apóstrofo inicial), Descripción, Débito, Crédito, Saldo.
  - Orden INVERSO: la fila 2 es "Saldo Total" (saldo final), luego los movimientos del más
    reciente al más antiguo, y la última fila es "Saldo Inicial". Se invierten para leer
    en orden cronológico.
  - El banco no trae totales de débitos/créditos ni número de operaciones: el cuadre se
    apoya en el saldo línea por línea y en saldo inicial → saldo final.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path
from typing import Sequence

import openpyxl
import xlrd

from ..modelo import Extracto, Movimiento, ResultadoLectura, TotalesBanco
from ..montos import CERO, MontoInvalido, monto_excel, monto_ve

BANCO = "BANPLUS"


class LectorBanplus:
    banco = BANCO

    def reconoce(self, ruta: Path) -> bool:
        if Path(ruta).suffix.lower() == ".xls":
            try:
                sh = xlrd.open_workbook(str(ruta)).sheet_by_index(0)
            except Exception:
                return False
            return [str(c).strip() for c in sh.row_values(0)[:3]] == ["Fecha", "Referencia", "Descripción/Concepto"]
        if Path(ruta).suffix.lower() != ".xlsx":
            return False
        ws = openpyxl.load_workbook(str(ruta), read_only=True).worksheets[0]
        return ws.title.startswith("Cuenta-0174")

    def leer(self, rutas: Sequence[Path]) -> ResultadoLectura:
        return ResultadoLectura(extractos=[self._leer_uno(Path(r)) for r in rutas])

    def _leer_uno(self, ruta: Path) -> Extracto:
        if ruta.suffix.lower() == ".xls":
            return self._leer_xls(ruta)
        ws = openpyxl.load_workbook(str(ruta), data_only=True).worksheets[0]
        filas = list(ws.iter_rows(values_only=True))
        enc = [str(c).strip() if c is not None else "" for c in filas[0]]
        esperado = ["Fecha", "Referencia", "Descripción", "Débito", "Crédito", "Saldo"]
        if enc[:6] != esperado:
            raise ValueError(f"{ruta.name}: encabezado inesperado {enc}")
        mc = re.search(r"Cuenta-(\d+)", ws.title)
        cuenta = mc.group(1) if mc else ""

        errores: list[str] = []
        saldo_ini = saldo_fin = None
        movs: list[Movimiento] = []
        for n, f in enumerate(filas[1:], 2):
            fecha, ref, desc, deb, cre, saldo = (list(f) + [None] * 6)[:6]
            if desc == "Saldo Total":
                saldo_fin = monto_excel(saldo)
                continue
            if desc == "Saldo Inicial":
                saldo_ini = monto_excel(saldo)
                continue
            if not isinstance(fecha, dt.datetime):
                if any(c not in (None, "") for c in f):
                    errores.append(f"fila {n}: no reconocida → {f}")
                continue
            try:
                movs.append(Movimiento(
                    fecha_contable=fecha.date(), fecha_real=fecha.date(),
                    referencia=str(ref or "").lstrip("'").strip(),
                    descripcion=str(desc or "").strip(),
                    debito=monto_excel(deb) if deb not in (None, "") else CERO,
                    credito=monto_excel(cre) if cre not in (None, "") else CERO,
                    saldo=monto_excel(saldo),
                ))
            except (MontoInvalido, ValueError) as e:
                errores.append(f"fila {n}: {e}")
        movs.reverse()  # el banco los lista del más reciente al más antiguo

        if saldo_ini is None:
            errores.append("No se encontró la fila 'Saldo Inicial'")
        if saldo_fin is None:
            errores.append("No se encontró la fila 'Saldo Total'")
        fechas = [m.fecha_contable for m in movs]
        avisos: list[str] = []
        desde = min(fechas) if fechas else None
        hasta = max(fechas) if fechas else dt.date.today()
        # El export no dice el período; si es un mes calendario, se toma el mes completo. Si trae unos
        # pocos movimientos de los primeros días del mes siguiente (p. ej. intereses del 01/10), el período
        # es el mes con más movimientos y se avisa; esos movimientos se mantienen (forman parte del saldo).
        if fechas and desde.day <= 3:
            desde = desde.replace(day=1)
            from collections import Counter
            (anio, mes), _ = Counter((f.year, f.month) for f in fechas).most_common(1)[0]
            fin_mes = (dt.date(anio, mes, 28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
            fuera = [f for f in fechas if f > fin_mes]
            if fuera and (anio, mes) == (desde.year, desde.month):
                avisos.append(f"El archivo incluye {len(fuera)} movimiento(s) posteriores al período "
                               f"({', '.join(sorted({f'{f:%d/%m/%Y}' for f in fuera}))}); se incluyen porque "
                               "forman parte del saldo final del archivo.")
                hasta = fin_mes
            else:
                sig = (hasta.replace(day=28) + dt.timedelta(days=4))
                hasta = sig - dt.timedelta(days=sig.day)
        return Extracto(
            banco=BANCO, cuenta=cuenta, desde=desde, hasta=hasta,
            saldo_anterior=saldo_ini if saldo_ini is not None else CERO,
            movimientos=movs,
            totales_banco=TotalesBanco(saldo_anterior=saldo_ini, nuevo_saldo=saldo_fin),
            archivo=ruta.name, errores_lectura=errores, avisos_lectura=avisos,
        )

    def _leer_xls(self, ruta: Path) -> Extracto:
        """Variante .xls (CACAO, agosto 2026):
        - Columnas: Fecha (texto AAAA-MM-DD) | Referencia | Descripción/Concepto | Débito | Crédito | Saldo,
          montos como texto 1.234,56.
        - Orden cronológico; fila 2 = "Saldo Inicial"; no trae saldo final ni totales.
        - Cada movimiento ocupa dos filas: la segunda (sin fecha) trae a veces un detalle
          ("Liquidación Automatica") que se agrega al movimiento anterior.
        """
        sh = xlrd.open_workbook(str(ruta)).sheet_by_index(0)
        errores: list[str] = []
        saldo_ini = None
        movs: list[Movimiento] = []
        for r in range(1, sh.nrows):
            fecha, ref, desc, deb, cre, saldo = [str(c).strip() for c in (sh.row_values(r) + [""] * 6)[:6]]
            if desc == "Saldo Inicial":
                saldo_ini = monto_ve(saldo)
                continue
            if not fecha:
                if desc and movs:
                    movs[-1].detalle = (movs[-1].detalle + " " + desc).strip()
                elif any((ref, deb, cre, saldo)):
                    errores.append(f"fila {r + 1}: sin fecha → {sh.row_values(r)}")
                continue
            try:
                f = dt.datetime.strptime(fecha, "%Y-%m-%d").date()
                movs.append(Movimiento(fecha_contable=f, fecha_real=f, referencia=ref, descripcion=desc,
                                       debito=monto_ve(deb) if deb else CERO,
                                       credito=monto_ve(cre) if cre else CERO, saldo=monto_ve(saldo)))
            except (MontoInvalido, ValueError) as e:
                errores.append(f"fila {r + 1}: {e}")
        if saldo_ini is None:
            errores.append("No se encontró la fila 'Saldo Inicial'")
        fechas = [m.fecha_contable for m in movs]
        hasta = max(fechas) if fechas else dt.date.today()
        sig = hasta.replace(day=28) + dt.timedelta(days=4)
        # Sin saldo final impreso: el único dato del banco para cerrar es el saldo de la última línea.
        return Extracto(
            banco=BANCO, cuenta="", desde=hasta.replace(day=1), hasta=sig - dt.timedelta(days=sig.day),
            saldo_anterior=saldo_ini if saldo_ini is not None else CERO, movimientos=movs,
            totales_banco=TotalesBanco(saldo_anterior=saldo_ini,
                                       nuevo_saldo=movs[-1].saldo if movs else None),
            archivo=ruta.name, errores_lectura=errores)

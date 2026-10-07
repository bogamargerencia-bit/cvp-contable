"""Ciclo de revisión: la app emite pendientes, el analista decide y reintroduce el archivo.

Corrida 1: `pendientes(rep)` lista lo que hay que revisar, cada partida con un CÓDIGO ESTABLE
           (se calcula del contenido, no de la posición), y el Excel lleva la hoja «Revisión»
           con tres columnas editables: Decisión, Comentario, Revisado por.
Corrida N: `leer_revision(excel_anterior)` lee esas decisiones; `aplicar(...)` las cruza con los
           pendientes de la nueva corrida (que puede usar un export del sistema ya corregido):
             - el código ya no aparece          → Resuelto
             - Aceptar / Justificado válidos     → Cerrado (la decisión se arrastra a la corrida siguiente)
             - «Corregido en el sistema» y sigue → Advertencia: se marcó corregido pero sigue igual
             - decisión incompleta               → abierto (Justificado sin comentario, sin revisor...)
             - código nuevo                      → Nuevo
           OK general = todos los bancos cuadran y no queda ninguna partida abierta.
           El cierre lo da el analista (la app registrará quién y cuándo en la Fase 2).
"""
from __future__ import annotations

import datetime as dt
import hashlib
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal
from enum import Enum
from pathlib import Path
from typing import Optional

import openpyxl

from .conciliacion import EstadoPartida as EP
from .conciliacion_caja import EstadoCaja
from .cuadre import Estado
from .diagnostico import (Explicacion, diagnosticar, explicar_caja_dia, explicar_diferencia_monto, explicar_divisa_fecha, explicar_grupo,
                          explicar_libro_sin_entradas, explicar_mov_suelto, explicar_otro_banco, explicar_pareja,
                          explicar_solo_libro, generico)
from .naturaleza import E_TRANSF, E_TRASLADO, OTROS, S_IMPUESTOS, S_PAGO_MOVIL, S_TRANSF, S_TRASLADO
from .proceso import ReporteCliente

VERSION = "cvp-revision-1"
TIPO_CORRECCION = "Caja: corrección de fecha propuesta"
ABREV = {"100_BANCO": "100", "BNC": "BNC", "BANPLUS": "BPL", "PLAZA": "PLZ", "BANPLUS+PLAZA": "BPZ"}
NOMBRE = {"100_BANCO": "100%", "BNC": "BNC", "BANPLUS": "Banplus", "PLAZA": "Plaza"}


class Decision(str, Enum):
    ACEPTAR = "Aceptar"
    JUSTIFICADO = "Justificado"
    CORREGIDO = "Corregido en el sistema"


DECISIONES = [d.value for d in Decision]


class Situacion(str, Enum):
    NUEVO = "Nuevo"
    SIN_DECISION = "Sin decisión"
    CERRADO = "Cerrado"
    CORREGIDO_SIGUE = "Marcado como corregido, pero sigue igual"
    INCOMPLETA = "Decisión incompleta"


ABIERTAS = {Situacion.NUEVO, Situacion.SIN_DECISION, Situacion.CORREGIDO_SIGUE, Situacion.INCOMPLETA}


@dataclass
class Pendiente:
    codigo: str
    banco: str
    origen: str                 # Banco / Libro / Caja / Grupo
    tipo: str                   # qué pasa (estado)
    fecha: Optional[dt.date]
    descripcion: str
    monto_bs: Optional[Decimal]
    monto_usd: Optional[Decimal]
    detalle: str
    sugerencia: str             # qué decisión suele corresponder
    propuesta: Optional[tuple[int, dt.date]] = None   # (fila de caja, fecha propuesta) si es una corrección
    situacion: Situacion = Situacion.NUEVO
    decision: str = ""
    comentario: str = ""
    revisado_por: str = ""
    aviso: str = ""             # por qué sigue abierta
    explicacion: str = ""       # qué pasó, en lenguaje claro (diagnostico.py)
    que_hacer: str = ""         # acción concreta para el analista


@dataclass
class DecisionLeida:
    codigo: str
    decision: str
    comentario: str
    revisado_por: str
    tipo: str
    descripcion: str


@dataclass
class RevisionAnterior:
    cliente: str
    periodo: str
    corrida: int
    decisiones: dict[str, DecisionLeida]
    historial: list[list]
    archivo: str
    propuestas: dict[str, tuple[int, dt.date]] = field(default_factory=dict)   # código → (fila, fecha)
    correcciones_previas: dict[int, dt.date] = field(default_factory=dict)     # ya aplicadas antes


@dataclass
class ResultadoRevision:
    corrida: int
    generado: dt.datetime
    pendientes: list[Pendiente]
    resueltas: list[DecisionLeida]           # estaban en la corrida anterior y ya no aparecen
    historial: list[list]
    ok_general: bool
    bancos_no_cuadran: list[str]
    advertencias: list[str] = field(default_factory=list)

    @property
    def abiertas(self) -> list[Pendiente]:
        return [p for p in self.pendientes if p.situacion in ABIERTAS]


# ---------------------------------------------------------------- generar pendientes
def _norm(t: str) -> str:
    t = unicodedata.normalize("NFKD", t or "").encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().upper()


def _codigo(banco: str, origen: str, partes: tuple, usados: dict[str, int]) -> str:
    h = hashlib.sha1("|".join(_norm(str(p)) for p in partes).encode()).hexdigest()[:6].upper()
    abrev = ABREV.get(banco) or (re.sub(r"[^A-Za-z0-9]", "", _norm(banco))[:3] or "XXX")
    base = f"{abrev}-{origen[0]}-{h}"
    usados[base] += 1
    return base if usados[base] == 1 else f"{base}-{usados[base]}"


# Naturalezas que se listan movimiento por movimiento (hay que identificar cada uno); el resto va agrupado.
INDIVIDUALES = {S_TRANSF, S_PAGO_MOVIL, S_TRASLADO, E_TRASLADO, E_TRANSF, S_IMPUESTOS, OTROS}
MAX_INDIVIDUALES = 25           # si hay más, se agrupan (con la lista de movimientos en la explicación)


def pendientes(rep: ReporteCliente) -> list[Pendiente]:
    out: list[Pendiente] = []
    usados: dict[str, int] = defaultdict(int)
    dg = diagnosticar(rep)

    def agregar(banco, origen, tipo, fecha, desc, bs, usd, detalle, sugerencia, clave_extra=(),
                exp: Optional[Explicacion] = None):
        # El código depende solo del contenido de la partida, no de la explicación.
        cod = _codigo(banco, origen, (tipo, banco, fecha.isoformat() if fecha else "", bs, usd, desc,
                                      *clave_extra), usados)
        exp = exp or generico(tipo)
        out.append(Pendiente(cod, banco, origen, tipo, fecha, desc, bs, usd, detalle, exp.decision or sugerencia,
                             explicacion=exp.texto, que_hacer=exp.que_hacer))

    for b in rep.bancos:
        c = rep.cuadres[b]
        # 1. El estado de cuenta no cuadra
        for d in c.diferencias:
            agregar(b, "Banco", "Estado de cuenta no cuadra", c.extracto.hasta, d.mensaje,
                    d.diferencia, None, "Revisar el PDF/archivo del banco o el lector.", "Corregido en el sistema",
                    exp=dg.cuadre.get(b))
        # 2. El export del sistema no es consistente
        for d in rep.libros[b].diferencias_saldo:
            agregar(b, "Libro", "Saldo del libro inconsistente", None, d, None, None,
                    "El saldo corrido del export no coincide con sus propios movimientos.", "Corregido en el sistema")

        rc = rep.conciliaciones[b]
        movs = rc.extracto.movimientos
        for p in rc.partidas:
            pareja = " + ".join(f"{movs[i].fecha_contable:%d/%m} {movs[i].descripcion} "
                                f"{movs[i].debito or movs[i].credito}" for i in p.movs)
            det = " · ".join(t for t in (f"Banco: {pareja}" if pareja else "", p.nota) if t)
            for j in p.asientos:
                a = rc.libro.asientos[j]
                desc = a.descripcion if a.monto_bs is not None else f"{a.referencia} — {a.descripcion}"
                if p.estado is EP.SOLO_LIBRO:
                    pr = dg.parejas.get((b, j))
                    agregar(b, "Libro", "Solo en libro", a.fecha, desc, a.monto_bs, a.monto_usd, det,
                            "Corregido en el sistema / Justificado (en tránsito)",
                            exp=explicar_pareja(rep, pr) if pr else explicar_solo_libro(rep, b, j, dg))
                elif p.estado is EP.OTRO_BANCO:
                    agregar(b, "Libro", "Registrado en el libro de otro banco", a.fecha, desc, a.monto_bs,
                            a.monto_usd, det, "Corregido en el sistema", exp=explicar_otro_banco(p.nota or ""))
                elif p.estado is EP.DIFERENCIA_MONTO:
                    agregar(b, "Libro", "Diferencia de monto", a.fecha, desc, a.monto_bs, a.monto_usd,
                            f"Diferencia Bs. {p.diferencia}. {det}", "Aceptar / Corregido en el sistema",
                            exp=explicar_diferencia_monto(p.diferencia))
                elif p.estado is EP.SENTIDO_INVERTIDO:
                    agregar(b, "Libro", "Sentido invertido", a.fecha, desc, a.monto_bs, a.monto_usd, det,
                            "Corregido en el sistema")
                elif p.estado is EP.RESUMEN_SIN_MONTO:
                    agregar(b, "Libro", "Asiento resumen sin conciliar", a.fecha, desc, None, a.monto_usd, det,
                            "Justificado")
                elif p.estado is EP.CONCILIADO_CAJA and "Faltan en el libro" in p.nota:
                    agregar(b, "Libro", "Asiento resumen incompleto", a.fecha, desc, None, a.monto_usd, p.nota,
                            "Corregido en el sistema")

    # 3. Cierre de caja
    sin_caja_idx: set[tuple[str, int]] = set()   # movimientos ya listados por la caja
    if rep.caja:
        for pr in rep.caja.propuestas:
            agregar("CAJA", "Caja", TIPO_CORRECCION, pr.fecha_actual,
                    f"Fila {pr.fila} ({pr.medio}): {pr.fecha_actual:%d/%m/%Y} → {pr.fecha_propuesta:%d/%m/%Y}",
                    pr.monto, None, pr.motivo, "Aceptar (la app corrige la fecha en la próxima corrida)")
            out[-1].propuesta = (pr.fila, pr.fecha_propuesta)
        for err in rep.caja.cierre.errores:
            agregar("CAJA", "Caja", "Cierre de caja con errores", None, err, None, None,
                    f"Archivo {rep.caja.cierre.archivo}", "Corregido (subir el cierre corregido) / Justificado")
        for t in rep.caja.totales_mes:
            agregar("+".join(t.bancos), "Caja", "Tarjetas: por abonar en el mes siguiente", None,
                    " + ".join(t.medios), t.por_abonar, None, t.nota,
                    "Justificado (verificar con el estado de cuenta del mes siguiente)")
        for l in rep.caja.lineas:
            sin_caja_idx.update(l.movs)
        for l in rep.caja.lineas:
            if l.estado in (EstadoCaja.POSIBLE_ERROR, EstadoCaja.NO_ENCONTRADO, EstadoCaja.EN_TRANSITO,
                            EstadoCaja.DIFERENCIA_DIA):
                sug = {EstadoCaja.POSIBLE_ERROR: "Aceptar / Corregido en el sistema",
                       EstadoCaja.NO_ENCONTRADO: "Corregido en el sistema / Justificado",
                       EstadoCaja.EN_TRANSITO: "Justificado (se verifica en el mes siguiente)",
                       EstadoCaja.DIFERENCIA_DIA: "Justificado / Corregido en la caja"}[l.estado]
                agregar(l.banco or "CAJA", "Caja", f"Caja: {l.estado.value}", l.fecha, l.medio, l.monto, None,
                        l.nota, sug,
                        exp=explicar_caja_dia(rep, dg, l) if l.estado is EstadoCaja.DIFERENCIA_DIA else None)
        lotes: dict[tuple[str, str], list[int]] = defaultdict(list)
        for b, i, motivo in rep.caja.sin_caja:
            sin_caja_idx.add((b, i))
            m = rep.conciliaciones[b].extracto.movimientos[i]
            if motivo.startswith(("Lote", "Abono POS de ventas del mes anterior")) or "sin día en el cierre" in motivo:
                clave = "Abono POS de ventas del mes anterior" if motivo.startswith("Abono POS") else motivo
                lotes[(b, clave)].append(i)
            else:
                agregar(b, "Banco", "Cobro sin registro en caja", m.fecha_contable,
                        f"{m.descripcion} {m.referencia}".strip(), m.credito, None, motivo,
                        "Justificado / Corregido en el sistema")
        for (b, motivo), idx in lotes.items():
            ms = [rep.conciliaciones[b].extracto.movimientos[i] for i in idx]
            agregar(b, "Banco", "Cobros sin día de caja" if "pago móvil" in motivo else "Lote sin día de caja",
                    ms[0].fecha_contable, motivo, sum((m.credito for m in ms), Decimal("0.00")), None,
                    f"{len(ms)} movimientos", "Justificado (ventas del mes anterior)" if "anterior" in motivo
                    else "Justificado / Corregido en la caja",
                    exp=generico("Abono POS de ventas del mes anterior") if "anterior" in motivo else None)

    # 3b. Cuentas en divisas (efectivo $, Zelle, USDT, fondo)
    for r in rep.divisas:
        n = r.cuenta.nombre
        if not r.ventas_ok:
            agregar(n, "Libro", "Divisas: ventas del mes ≠ asiento de ventas", None, n, None, r.ventas_mes,
                    r.ventas_nota, "Corregido en el sistema")
        for d in r.dias_con_diferencia:
            agregar(n, "Caja", "Divisas: ventas del día ≠ Kardex", d.fecha, n, None, d.diferencia,
                    f"Ventas US$ {d.ventas} · Kardex US$ {d.kardex}. {d.nota}".strip(), "Justificado / Corregido")
        # Pares con el mismo monto fuera de ±días: se explican juntos (fila del Kardex → asiento).
        pares = {p.pista.fila: r.libro.asientos[p.asientos[0]] for p in r.partidas if p.pista and p.asientos}
        for p in r.partidas:
            if p.estado in ("Conciliado", "Conciliado (agrupado)"):
                continue
            if p.asientos:
                a = r.libro.asientos[p.asientos[0]]
                agregar(n, "Libro", f"Divisas: {p.estado.lower()}", a.fecha, a.descripcion, None,
                        a.debito_usd - a.credito_usd, p.nota,
                        "Aceptar" if p.estado.startswith("Incluido") else "Corregido / Justificado",
                        exp=explicar_divisa_fecha(n, a, p.pista, "libro") if p.pista else None)
            for k in p.kardex:
                a = pares.get(k.fila)
                agregar(n, "Caja", "Divisas: solo en Kardex", k.fecha, k.descripcion, None, k.monto,
                        f"Kardex fila {k.fila}" + (f" · posible: libro {a.fecha:%d/%m} {a.descripcion}" if a else ""),
                        "Corregido en el sistema / Justificado",
                        exp=explicar_divisa_fecha(n, a, k, "kardex") if a else None)
        for nombre_s, (lib, kx) in (("inicial", r.saldo_ini), ("final", r.saldo_fin)):
            if kx is not None and lib != kx:
                agregar(n, "Libro", f"Divisas: saldo {nombre_s} libro ≠ Kardex", None, n, None, lib - kx,
                        f"Libro US$ {lib} · Kardex US$ {kx}", "Corregido / Justificado")
    for txt in rep.cruces_divisas_bancos:
        if "diferencia US$ 0.00" not in txt:
            agregar("DIVISAS", "Libro", "Venta de divisas: el banco la registra a otra tasa", None, txt[:120],
                    None, None, txt, "Corregido en el sistema / Justificado (indicar la tasa usada)")

    # 4. Libro sin ninguna entrada en el mes (falta el asiento de ventas)
    for b in dg.libro_sin_entradas:
        agregar(b, "Libro", "Libro sin entradas en el mes", rep.conciliaciones[b].extracto.hasta,
                f"Libro {NOMBRE.get(b, b)} sin débitos", None, None, "", "Corregido en el sistema",
                exp=explicar_libro_sin_entradas(rep, b))

    # 5. Lo que queda solo en el banco: transferencias, pagos, traslados y no clasificados uno por uno
    #    (hay que identificar cada uno); comisiones, cargos, ISLR, nómina... agrupados por naturaleza.
    #    Los movimientos que el diagnóstico emparejó con un asiento ya se explican en la partida del asiento.
    for b in rep.bancos:
        rc = rep.conciliaciones[b]
        grupos: dict[tuple[str, str], list] = defaultdict(list)
        for x in rep.clasificados[b]:
            p = rc.estado_mov.get(x.indice)
            if (p and p.estado in (EP.SOLO_BANCO, EP.CAJA_SIN_LIBRO) and (b, x.indice) not in sin_caja_idx
                    and (b, x.indice) not in dg.movs_en_pareja):
                grupos[(p.estado.value, x.naturaleza)].append(x)
        for (estado, nat), xs in sorted(grupos.items()):
            if estado == EP.SOLO_BANCO.value and nat in INDIVIDUALES and len(xs) <= MAX_INDIVIDUALES:
                for x in xs:
                    m = x.mov
                    agregar(b, "Banco", f"{estado}: {nat}", m.fecha_contable,
                            f"{m.descripcion} {m.referencia}".strip(), m.credito - m.debito, None,
                            " · ".join(t for t in (m.detalle, x.contraparte) if t),
                            "Corregido en el sistema", exp=explicar_mov_suelto(rep, b, x, dg))
                continue
            ms = [x.mov for x in xs]
            ent = sum((m.credito for m in ms), Decimal("0.00"))
            sal = sum((m.debito for m in ms), Decimal("0.00"))
            agregar(b, "Grupo", f"{estado}: {nat}", None, f"{nat} ({len(ms)} mov.)", ent - sal, None,
                    f"Entradas Bs. {ent} · salidas Bs. {sal}. Detalle en la hoja Banco "
                    f"{NOMBRE.get(b, b)} (filtrar Naturaleza y Estado).",
                    "Justificado (indicar el asiento) / Corregido en el sistema", exp=explicar_grupo(nat, estado, ms))
    return out


# ---------------------------------------------------------------- leer el archivo reintroducido
def leer_revision(ruta: Path) -> RevisionAnterior:
    ruta = Path(ruta)
    wb = openpyxl.load_workbook(str(ruta), data_only=True)
    if "_control" not in wb.sheetnames or "Revisión" not in wb.sheetnames:
        raise ValueError(f"{ruta.name}: no es un archivo de revisión emitido por la app "
                         "(faltan las hojas «Revisión» o «_control»).")
    filas_ctl = [r for r in wb["_control"].iter_rows(values_only=True) if r and r[0]]
    ctl = {str(r[0]): r[1] for r in filas_ctl if r[0] not in ("codigo", "propuesta", "correccion")}

    def fecha(v) -> dt.date:
        return v.date() if isinstance(v, dt.datetime) else dt.date.fromisoformat(str(v))
    propuestas = {str(r[1]): (int(r[2]), fecha(r[3])) for r in filas_ctl if r[0] == "propuesta"}
    correcciones_previas = {int(r[1]): fecha(r[2]) for r in filas_ctl if r[0] == "correccion"}
    if ctl.get("version") != VERSION:
        raise ValueError(f"{ruta.name}: versión de archivo no reconocida ({ctl.get('version')}).")
    codigos_emitidos = {str(r[1]).strip() for r in filas_ctl if r[0] == "codigo"}

    ws = wb["Revisión"]
    fila_enc = next(r for r in range(1, 40) if ws.cell(r, 1).value == "Código")
    enc = {str(ws.cell(fila_enc, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(fila_enc, c).value}
    decisiones: dict[str, DecisionLeida] = {}
    for r in range(fila_enc + 1, ws.max_row + 1):
        cod = ws.cell(r, enc["Código"]).value
        if not cod:
            continue
        cod = str(cod).strip()
        if cod not in codigos_emitidos:
            continue  # fila agregada a mano o de otra sección: se ignora
        val = lambda h: str(ws.cell(r, enc[h]).value or "").strip()
        decisiones[cod] = DecisionLeida(cod, val("Decisión"), val("Comentario"), val("Revisado por"),
                                        val("Tipo"), val("Descripción"))
    historial = []
    if "Historial" in wb.sheetnames:
        h = wb["Historial"]
        historial = [list(row[:9]) for row in h.iter_rows(min_row=4, values_only=True)
                     if row and isinstance(row[0], (int, float)) and not isinstance(row[0], bool)]
    return RevisionAnterior(str(ctl.get("cliente")), str(ctl.get("periodo")), int(ctl.get("corrida", 1)),
                            decisiones, historial, ruta.name, propuestas, correcciones_previas)


def correcciones_caja(anterior: RevisionAnterior) -> dict[int, dt.date]:
    """Correcciones de fecha a aplicar en la nueva corrida: las ya aplicadas antes más las propuestas
    que el analista aceptó (con decisión válida) en el archivo reintroducido."""
    out = dict(anterior.correcciones_previas)
    for cod, (fila, fecha) in anterior.propuestas.items():
        d = anterior.decisiones.get(cod)
        if d and d.decision == Decision.ACEPTAR.value and _validar(d) is None:
            out[fila] = fecha
    return out


# ---------------------------------------------------------------- aplicar decisiones
def _validar(d: DecisionLeida) -> Optional[str]:
    if d.decision not in DECISIONES:
        return f"Decisión «{d.decision}» no válida (use: {', '.join(DECISIONES)})."
    if not d.revisado_por:
        return "Falta «Revisado por»."
    if d.decision == Decision.JUSTIFICADO.value and not d.comentario:
        return "«Justificado» requiere un comentario."
    return None


def aplicar(rep: ReporteCliente, anterior: Optional[RevisionAnterior] = None,
            ahora: Optional[dt.datetime] = None) -> ResultadoRevision:
    ahora = ahora or dt.datetime.now().replace(microsecond=0)
    lista = pendientes(rep)
    advertencias: list[str] = []
    historial: list[list] = []
    resueltas: list[DecisionLeida] = []
    corrida = 1
    if anterior is not None:
        if (anterior.cliente, anterior.periodo) != (rep.cliente, rep.periodo):
            raise ValueError(f"El archivo de revisión es de {anterior.cliente} {anterior.periodo}, "
                             f"no de {rep.cliente} {rep.periodo}.")
        corrida = anterior.corrida + 1
        historial = list(anterior.historial)
        # Última entrada registrada por código, para no repetir la misma decisión en cada corrida.
        ultima = {fila[2]: tuple(fila[5:9]) for fila in historial}

        def anotar(fila: list) -> None:
            if ultima.get(fila[2]) != tuple(fila[5:9]):
                historial.append(fila)
                ultima[fila[2]] = tuple(fila[5:9])
        actuales = {p.codigo: p for p in lista}
        for cod, d in anterior.decisiones.items():
            if cod not in actuales:
                resueltas.append(d)
                anotar([anterior.corrida, ahora, cod, d.tipo, d.descripcion, d.decision or "—",
                        d.comentario or None, d.revisado_por or None, "Resuelto: ya no aparece en la corrida "
                        f"{corrida}"])
        for p in lista:
            d = anterior.decisiones.get(p.codigo)
            if d is None:
                p.situacion = Situacion.NUEVO
                continue
            p.decision, p.comentario, p.revisado_por = d.decision, d.comentario, d.revisado_por
            if not d.decision:
                p.situacion = Situacion.SIN_DECISION
                continue
            error = _validar(d)
            if error:
                p.situacion, p.aviso = Situacion.INCOMPLETA, error
            elif d.decision == Decision.CORREGIDO.value:
                p.situacion = Situacion.CORREGIDO_SIGUE
                p.aviso = ("Se marcó como corregido en el sistema, pero sigue apareciendo igual. "
                           "¿Se subió el export nuevo del sistema?")
            else:
                p.situacion = Situacion.CERRADO
            anotar([anterior.corrida, ahora, p.codigo, p.tipo, p.descripcion, d.decision,
                    d.comentario or None, d.revisado_por or None, p.situacion.value])
        if not anterior.decisiones:
            advertencias.append("El archivo reintroducido no trae ninguna decisión.")
    no_cuadran = [b for b in rep.bancos if rep.cuadres[b].estado is Estado.NO_CUADRA]
    ok = not no_cuadran and not any(p.situacion in ABIERTAS for p in lista)
    return ResultadoRevision(corrida, ahora, lista, resueltas, historial, ok, no_cuadran, advertencias)

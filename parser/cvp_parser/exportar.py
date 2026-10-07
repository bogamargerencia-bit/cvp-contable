"""Excel de conciliación de un cliente (todos sus bancos en un libro).

Hojas:
  Revisión           EDITABLE. Partidas a revisar con código estable; el analista decide y reintroduce el
                     archivo (ver revision.py). Dice si hay OK general.
  Resumen            Cuadre de cada banco, estado del libro y conciliación por estado.
  Naturaleza         Movimientos del banco por naturaleza (todos y los que no tienen asiento).
  Cierre de caja     (si se dio el Excel de ventas) caja ↔ banco ↔ asientos resumen.
  Historial          Decisiones de corridas anteriores y su resultado.
  _control           (oculta) versión, cliente, período, corrida y códigos emitidos.
  Asientos resumen   Asientos del libro sin monto en Bs. frente a las entradas del banco sin asiento.
  Banco <x>          Todos los movimientos del banco con naturaleza y estado de conciliación.
  Libro <x>          Todos los asientos del sistema con su estado y su pareja en el banco.

Los totales y controles son fórmulas sobre las hojas de detalle (se recalculan si se editan).
"""
from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

from .conciliacion import EstadoPartida as EP
from .conciliacion_caja import CONCILIADOS_CAJA, EstadoCaja
from .cuadre import Estado
from .naturaleza import E_PAGO_MOVIL, E_POS_CRE, E_POS_DEB, E_TRANSF, ENTRADAS, ORDEN

COBROS = {E_POS_DEB, E_POS_CRE, E_PAGO_MOVIL, E_TRANSF}
from .proceso import ReporteCliente

NOMBRE = {"100_BANCO": "100%", "BNC": "BNC", "BANPLUS": "Banplus", "PLAZA": "Plaza"}

F = Font(name="Arial", size=10)
B = Font(name="Arial", size=10, bold=True)
AZUL = Font(name="Arial", size=10, color="0000FF")
VERDE_TXT = Font(name="Arial", size=10, color="008000")
BLANCA = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TIT = Font(name="Arial", size=13, bold=True)
SUB = Font(name="Arial", size=11, bold=True)
NOTA = Font(name="Arial", size=9, italic=True, color="595959")
ROJA = Font(name="Arial", size=10, bold=True, color="C00000")
ENC = PatternFill("solid", start_color="1F3864")
TOT = PatternFill("solid", start_color="D9E1F2")
FILL = {
    EP.CONCILIADO: PatternFill("solid", start_color="E2EFDA"),
    EP.CONCILIADO_AGRUPADO: PatternFill("solid", start_color="C6EFCE"),
    EP.CONCILIADO_COMPONENTES: PatternFill("solid", start_color="C6EFCE"),
    EP.DIFERENCIA_MONTO: PatternFill("solid", start_color="FFEB9C"),
    EP.SENTIDO_INVERTIDO: PatternFill("solid", start_color="FFEB9C"),
    EP.OTRO_BANCO: PatternFill("solid", start_color="FCE4D6"),
    EP.SOLO_LIBRO: PatternFill("solid", start_color="FFC7CE"),
    EP.RESUMEN_SIN_MONTO: PatternFill("solid", start_color="DDEBF7"),
    EP.CONCILIADO_CAJA: PatternFill("solid", start_color="D8E4BC"),
    EP.CONCILIADO_REFERENCIA: PatternFill("solid", start_color="E2EFDA"),
    EP.CONCILIADO_TOTAL: PatternFill("solid", start_color="C6EFCE"),
    EP.CAJA_SIN_LIBRO: PatternFill("solid", start_color="FCE4D6"),
}
FILL_CAJA = {
    EstadoCaja.CONCILIADO: PatternFill("solid", start_color="E2EFDA"),
    EstadoCaja.AGRUPADO: PatternFill("solid", start_color="C6EFCE"),
    EstadoCaja.LOTES: PatternFill("solid", start_color="C6EFCE"),
    EstadoCaja.POSIBLE_ERROR: PatternFill("solid", start_color="FFEB9C"),
    EstadoCaja.EN_TRANSITO: PatternFill("solid", start_color="DDEBF7"),
    EstadoCaja.NO_ENCONTRADO: PatternFill("solid", start_color="FFC7CE"),
    EstadoCaja.TOTAL_MES: PatternFill("solid", start_color="EDEDED"),
    EstadoCaja.DIFERENCIA_DIA: PatternFill("solid", start_color="FFEB9C"),
}
LINEA = Border(bottom=Side(style="thin", color="BFBFBF"))
NUM = '#,##0.00;-#,##0.00;"-"'
TASA = "#,##0.0000"
FECHA = "DD/MM/YYYY"
ORDEN_ESTADOS = [EP.CONCILIADO, EP.CONCILIADO_REFERENCIA, EP.CONCILIADO_AGRUPADO, EP.CONCILIADO_COMPONENTES,
                 EP.CONCILIADO_TOTAL, EP.CONCILIADO_CAJA,
                 EP.DIFERENCIA_MONTO, EP.SENTIDO_INVERTIDO, EP.OTRO_BANCO, EP.CAJA_SIN_LIBRO, EP.SOLO_LIBRO,
                 EP.RESUMEN_SIN_MONTO, EP.SOLO_BANCO]


def _q(hoja: str) -> str:
    return "'" + hoja.replace("'", "''") + "'"


def _enc(ws, fila: int, titulos: list[str], anchos: list[float] | None = None) -> None:
    for i, t in enumerate(titulos, 1):
        c = ws.cell(fila, i, t)
        c.font, c.fill = BLANCA, ENC
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if anchos:
            ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1]
    ws.row_dimensions[fila].height = 30


def _fila(ws, r: int, valores: list, fmt: dict[int, str] | None = None, fill=None, font=F) -> None:
    for i, v in enumerate(valores, 1):
        c = ws.cell(r, i, v)
        c.font = font
        c.border = LINEA
        if fmt and i in fmt:
            c.number_format = fmt[i]
        if fill is not None:
            c.fill = fill


def excel_conciliacion(rep: ReporteCliente, salida: Path, revision=None) -> Path:
    """revision: ResultadoRevision de revision.aplicar(). Si no se pasa, se calcula como primera corrida."""
    from .revision import aplicar
    if revision is None:
        revision = aplicar(rep)
    wb = Workbook()
    hoja_rev = wb.active
    hoja_rev.title = "Revisión"
    resumen = wb.create_sheet("Resumen")
    hoja_nat = wb.create_sheet("Naturaleza")
    hoja_caja = wb.create_sheet("Cierre de caja") if rep.caja else None
    quedan_resumen = any(p.estado is EP.RESUMEN_SIN_MONTO for r in rep.conciliaciones.values() for p in r.partidas)
    hoja_res = wb.create_sheet("Asientos resumen") if quedan_resumen else None

    rangos: dict[str, dict[str, str]] = {}
    for b in rep.bancos:
        rangos[b] = _hoja_banco(wb, rep, b)
        rangos[b].update(_hoja_libro(wb, rep, b))

    _hoja_resumen(resumen, rep, rangos)
    _hoja_naturaleza(hoja_nat, rep, rangos)
    _hoja_revision(hoja_rev, rep, revision)
    _hoja_historial(wb.create_sheet("Historial"), rep, revision)
    _hoja_control(wb.create_sheet("_control"), rep, revision)
    if hoja_caja is not None:
        _hoja_caja(hoja_caja, rep)
    if rep.divisas:
        _hoja_divisas(wb.create_sheet("Divisas", wb.sheetnames.index("Naturaleza") + 1), rep)
    if hoja_res is not None:
        _hoja_asientos_resumen(hoja_res, rep, rangos)
    salida = Path(salida)
    wb.save(salida)
    return salida


# ------------------------------------------------------------------ Banco <x>
def _hoja_banco(wb, rep: ReporteCliente, b: str) -> dict[str, str]:
    nombre = f"Banco {NOMBRE.get(b, b)}"
    ws = wb.create_sheet(nombre)
    e = rep.cuadres[b].extracto
    r_conc = rep.conciliaciones[b]
    libro = rep.libros[b]
    ws["A1"] = f"{e.titular or b} — {NOMBRE.get(b, b)} cuenta {e.cuenta}"
    ws["A1"].font = TIT
    ws["A2"] = f"Período {e.desde:%d/%m/%Y} al {e.hasta:%d/%m/%Y} · Archivo {e.archivo} · Saldo inicial (banco):"
    ws["A2"].font = NOTA
    ws["H2"] = e.saldo_anterior
    ws["H2"].font, ws["H2"].number_format = AZUL, NUM
    _enc(ws, 4, ["Fecha real", "Fecha contable", "Referencia", "Descripción", "Detalle",
                 "Débito (Bs.)", "Crédito (Bs.)", "Saldo banco (Bs.)", "Saldo calculado (Bs.)",
                 "Naturaleza", "Estado conciliación", "Fila libro", "Fecha libro", "Descripción libro",
                 "Monto libro (Bs.)", "Diferencia (Bs.)", "Nota"],
         [11, 11, 15, 30, 40, 15, 15, 16, 16, 30, 24, 8, 11, 34, 15, 12, 50])
    r = 5
    for x in rep.clasificados[b]:
        m = x.mov
        p = r_conc.estado_mov.get(x.indice)
        asis = [libro.asientos[j] for j in p.asientos] if p else []
        prev = "$H$2" if r == 5 else f"I{r - 1}"
        nota = " · ".join(t for t in (p.nota if p else "", x.contraparte) if t)
        _fila(ws, r, [
            m.fecha_real, m.fecha_contable, m.referencia or None, m.descripcion, m.detalle or None,
            m.debito or None, m.credito or None, m.saldo,
            f"={prev}-F{r}+G{r}",
            x.naturaleza, p.estado.value if p else "",
            ", ".join(str(a.fila) for a in asis) or None,
            asis[0].fecha if asis else None,
            " + ".join(a.descripcion for a in asis) or None,
            sum((a.monto_bs or Decimal(0) for a in asis), Decimal("0.00")) if asis else None,
            f'=IF(K{r}="{EP.DIFERENCIA_MONTO.value}",O{r}-F{r}-G{r},"")',
            nota or None,
        ], {1: FECHA, 2: FECHA, 6: NUM, 7: NUM, 8: NUM, 9: NUM, 13: FECHA, 15: NUM, 16: NUM},
            fill=FILL.get(p.estado) if p else None)
        r += 1
    ultima = r - 1
    ws.cell(r, 4, "TOTALES").font = B
    for col in "FG":
        c = ws[f"{col}{r}"]
        c.value = f"=SUM({col}5:{col}{ultima})"
        c.font, c.number_format, c.fill = B, NUM, TOT
    ws.cell(r + 1, 4, "Control: saldo calculado − saldo banco (debe ser 0)").font = NOTA
    c = ws.cell(r + 1, 9, f"=ROUND(I{ultima}-H{ultima},2)")
    c.font, c.number_format = B, NUM
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:Q{ultima}"
    h = _q(nombre)
    return {"banco_hoja": nombre, "deb": f"{h}!$F$5:$F${ultima}", "cre": f"{h}!$G$5:$G${ultima}",
            "nat": f"{h}!$J$5:$J${ultima}", "est_b": f"{h}!$K$5:$K${ultima}",
            "tot_deb": f"{h}!$F${r}", "tot_cre": f"{h}!$G${r}", "saldo_ini": f"{h}!$H$2",
            "saldo_fin_calc": f"{h}!$I${ultima}", "control": f"{h}!$I${r + 1}"}


# ------------------------------------------------------------------ Libro <x>
def _hoja_libro(wb, rep: ReporteCliente, b: str) -> dict[str, str]:
    nombre = f"Libro {NOMBRE.get(b, b)}"
    ws = wb.create_sheet(nombre)
    libro = rep.libros[b]
    r_conc = rep.conciliaciones[b]
    movs = rep.cuadres[b].extracto.movimientos
    ws["A1"] = f"Libro de bancos del sistema — {NOMBRE.get(b, b)} ({libro.archivo})"
    ws["A1"].font = TIT
    ws["A2"] = "Montos del sistema en US$; la Referencia es el monto en Bs. Saldo inicial US$ (deducido de la 1ª línea):"
    ws["A2"].font = NOTA
    ws["H2"] = libro.saldo_inicial_usd
    ws["H2"].font, ws["H2"].number_format = AZUL, NUM
    _enc(ws, 4, ["Fila", "Fecha", "Comprobante", "Referencia", "Descripción", "Débitos (US$)",
                 "Créditos (US$)", "Saldo libro (US$)", "Saldo calculado (US$)", "Monto (Bs.)",
                 "Tasa implícita (Bs./US$)", "Estado conciliación", "Fecha banco", "Descripción banco",
                 "Monto banco (Bs.)", "Nota"],
         [6, 11, 16, 13, 40, 13, 13, 15, 15, 15, 12, 24, 11, 34, 15, 60])
    r = 5
    for j, a in enumerate(libro.asientos):
        p = r_conc.estado_asiento.get(j)
        ms = [movs[i] for i in p.movs] if p else []
        prev = "$H$2" if r == 5 else f"I{r - 1}"
        _fila(ws, r, [
            a.fila, a.fecha, a.comprobante, a.referencia, a.descripcion,
            a.debito_usd or None, a.credito_usd or None, a.saldo_usd,
            f"={prev}+F{r}-G{r}",
            a.monto_bs, f'=IF(J{r}="","",IFERROR(J{r}/(F{r}+G{r}),""))',
            p.estado.value if p else "",
            ms[0].fecha_contable if ms else None,
            " + ".join(m.descripcion for m in ms) or None,
            sum((m.debito or m.credito for m in ms), Decimal("0.00")) if ms else None,
            p.nota if p and p.nota else None,
        ], {2: FECHA, 6: NUM, 7: NUM, 8: NUM, 9: NUM, 10: NUM, 11: TASA, 13: FECHA, 15: NUM},
            fill=FILL.get(p.estado) if p else None)
        r += 1
    ultima = r - 1
    ws.cell(r, 5, "TOTALES").font = B
    for col in "FGJ":
        c = ws[f"{col}{r}"]
        c.value = f"=SUM({col}5:{col}{ultima})"
        c.font, c.number_format, c.fill = B, NUM, TOT
    ws.cell(r + 1, 5, "Control: saldo calculado − saldo del libro en la última línea (debe ser 0)").font = NOTA
    c = ws.cell(r + 1, 9, f"=ROUND(I{ultima}-H{ultima},2)")
    c.font, c.number_format = B, NUM
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:P{ultima}"
    h = _q(nombre)
    return {"libro_hoja": nombre, "bs_l": f"{h}!$J$5:$J${ultima}", "est_l": f"{h}!$L$5:$L${ultima}",
            "tasa": f"{h}!$K$5:$K${ultima}", "l_ini": f"{h}!$H$2", "l_fin": f"{h}!$H${ultima}",
            "l_fin_calc": f"{h}!$I${ultima}", "l_control": f"{h}!$I${r + 1}"}


# ------------------------------------------------------------------ Resumen
def _hoja_resumen(ws, rep: ReporteCliente, rg: dict[str, dict[str, str]]) -> None:
    ws["A1"] = f"{rep.cliente} — Cuadre y conciliación bancaria {rep.periodo}"
    ws["A1"].font = TIT
    ws["A2"] = ("Valores en azul: tomados de los archivos. El resto son fórmulas sobre las hojas de detalle. "
                "Montos del banco en Bs.; el libro del sistema está en US$ con el monto en Bs. en la Referencia.")
    ws["A2"].font = NOTA
    for col, w in zip("ABCDEFG", [44, 18, 18, 18, 18, 18, 18]):
        ws.column_dimensions[col].width = w
    r = 4
    if rep.alertas:
        ws.cell(r, 1, "ATENCIÓN").font = ROJA
        r += 1
        for al in rep.alertas:
            ws.cell(r, 1, al).font = ROJA
            r += 1
        r += 1
    for b in rep.bancos:
        c = rep.cuadres[b]
        e = c.extracto
        g = rg[b]
        ws.cell(r, 1, f"{NOMBRE.get(b, b)} · {e.titular or '(el archivo del banco no trae el titular)'} · cuenta {e.cuenta}").font = SUB
        r += 1
        # Cuadre
        _enc(ws, r, ["Cuadre contra el banco", "Bs."])
        r += 1
        filas = [
            ("Saldo inicial", f"={g['saldo_ini']}"),
            ("(−) Débitos del detalle", f"={g['tot_deb']}"),
            ("(+) Créditos del detalle", f"={g['tot_cre']}"),
            ("Saldo final calculado", f"=ROUND(B{r}-B{r + 1}+B{r + 2},2)"),
            ("Saldo final según el banco", e.totales_banco.nuevo_saldo if e.totales_banco else None),
            ("Diferencia (debe ser 0)", f"=ROUND(B{r + 3}-B{r + 4},2)"),
        ]
        for k, (t, v) in enumerate(filas):
            ws.cell(r + k, 1, t).font = F
            cc = ws.cell(r + k, 2, v)
            cc.number_format = NUM
            cc.font = AZUL if k == 4 else (B if k == 5 else VERDE_TXT if isinstance(v, str) and "!" in v else F)
        r += len(filas)
        estado_txt = {Estado.CUADRA: "CUADRA al céntimo",
                      Estado.REQUIERE_REVISION: "CUADRA, con datos que el banco no informa (ver abajo)",
                      Estado.NO_CUADRA: "NO CUADRA"}[c.estado]
        ws.cell(r, 1, "Resultado del cuadre").font = B
        ws.cell(r, 2, estado_txt).font = ROJA if c.estado is Estado.NO_CUADRA else B
        r += 1
        for d in c.diferencias:
            ws.cell(r, 1, f"Diferencia: {d.mensaje}").font = ROJA
            r += 1
        for a in c.avisos:
            ws.cell(r, 1, f"Nota: {a.mensaje}").font = NOTA
            r += 1
        # Libro
        r += 1
        _enc(ws, r, ["Libro del sistema", "US$"])
        r += 1
        for t, v in [("Saldo inicial (deducido de la 1ª línea)", f"={g['l_ini']}"),
                     ("Saldo final según el libro", f"={g['l_fin']}"),
                     ("Saldo final recalculado", f"=ROUND({g['l_fin_calc']},2)"),
                     ("Diferencia en el propio libro (debe ser 0)", f"={g['l_control']}")]:
            ws.cell(r, 1, t).font = F
            cc = ws.cell(r, 2, v)
            cc.number_format, cc.font = NUM, VERDE_TXT
            r += 1
        for d in rep.libros[b].diferencias_saldo:
            ws.cell(r, 1, f"Diferencia en el libro: {d}").font = ROJA
            r += 1
        if rep.libros[b].filas_ignoradas:
            ws.cell(r, 1, "Filas del export ignoradas (sin fecha): "
                    + "; ".join(rep.libros[b].filas_ignoradas)).font = NOTA
            r += 1
        # Conciliación
        r += 1
        _enc(ws, r, ["Conciliación", "Mov. banco", "Monto banco (Bs.)", "Asientos libro", "Monto libro (Bs.)"])
        r += 1
        r0 = r
        for est in ORDEN_ESTADOS:
            ws.cell(r, 1, est.value).font = F
            ws.cell(r, 2, f'=COUNTIF({g["est_b"]},A{r})')
            ws.cell(r, 3, f'=SUMIFS({g["deb"]},{g["est_b"]},A{r})+SUMIFS({g["cre"]},{g["est_b"]},A{r})')
            ws.cell(r, 4, f'=COUNTIF({g["est_l"]},A{r})')
            ws.cell(r, 5, f'=SUMIFS({g["bs_l"]},{g["est_l"]},A{r})')
            for k in range(1, 6):
                ws.cell(r, k).border = LINEA
                if k > 1:
                    ws.cell(r, k).font = F
                    ws.cell(r, k).number_format = NUM if k in (3, 5) else "0"
                if est in FILL:
                    ws.cell(r, k).fill = FILL[est]
            r += 1
        ws.cell(r, 1, "TOTAL").font = B
        for k in range(2, 6):
            col = get_column_letter(k)
            cc = ws.cell(r, k, f"=SUM({col}{r0}:{col}{r - 1})")
            cc.font, cc.fill = B, TOT
            cc.number_format = NUM if k in (3, 5) else "0"
        ws.cell(r, 1).fill = TOT
        r += 3


# ------------------------------------------------------------------ Naturaleza
def _hoja_naturaleza(ws, rep: ReporteCliente, rg: dict[str, dict[str, str]]) -> None:
    ws["A1"] = f"{rep.cliente} — Movimientos bancarios por naturaleza ({rep.periodo})"
    ws["A1"].font = TIT
    ws["A2"] = ("Montos en Bs. tomados de las hojas Banco. La naturaleza sale de la descripción del banco; "
                "los traslados entre cuentas propias se detectan por monto y fecha entre bancos, y la nómina "
                "en 100% y Banplus se toma del asiento del libro cuando está conciliado.")
    ws["A2"].font = NOTA
    ws.column_dimensions["A"].width = 52
    usadas = [n for n in ORDEN if any(x.naturaleza == n for l in rep.clasificados.values() for x in l)]

    def tabla(r: int, titulo: str, solo_banco: bool) -> int:
        ws.cell(r, 1, titulo).font = SUB
        r += 1
        cols = ["Naturaleza"]
        for b in rep.bancos:
            cols += [f"{NOMBRE.get(b, b)} cant.", f"{NOMBRE.get(b, b)} entradas", f"{NOMBRE.get(b, b)} salidas"]
        cols += ["Total entradas", "Total salidas"]
        _enc(ws, r, cols)
        r += 1
        r0 = r
        extra = (lambda g: f',{g["est_b"]},"{EP.SOLO_BANCO.value}"') if solo_banco else (lambda g: "")
        for n in usadas:
            ws.cell(r, 1, n).font = B if n in ENTRADAS else F
            k = 2
            ent, sal = [], []
            for b in rep.bancos:
                g = rg[b]
                ws.cell(r, k, f'=COUNTIFS({g["nat"]},$A{r}{extra(g)})').number_format = "0"
                ws.cell(r, k + 1, f'=SUMIFS({g["cre"]},{g["nat"]},$A{r}{extra(g)})').number_format = NUM
                ws.cell(r, k + 2, f'=SUMIFS({g["deb"]},{g["nat"]},$A{r}{extra(g)})').number_format = NUM
                ent.append(f"{get_column_letter(k + 1)}{r}")
                sal.append(f"{get_column_letter(k + 2)}{r}")
                k += 3
            ws.cell(r, k, "=" + "+".join(ent)).number_format = NUM
            ws.cell(r, k + 1, "=" + "+".join(sal)).number_format = NUM
            for c in range(1, k + 2):
                ws.cell(r, c).border = LINEA
                if c > 1:
                    ws.cell(r, c).font = F
            r += 1
        ncols = 1 + 3 * len(rep.bancos) + 2
        ws.cell(r, 1, "TOTAL").font = B
        for c in range(2, ncols + 1):
            col = get_column_letter(c)
            cc = ws.cell(r, c, f"=SUM({col}{r0}:{col}{r - 1})")
            cc.font, cc.fill = B, TOT
            cc.number_format = "0" if (c - 2) % 3 == 0 and c <= 1 + 3 * len(rep.bancos) else NUM
        ws.cell(r, 1).fill = TOT
        if not solo_banco:
            r += 1
            ws.cell(r, 1, "Control: total según hoja Banco").font = NOTA
            r_ctrl = r
            k = 2
            for b in rep.bancos:
                g = rg[b]
                ws.cell(r, k + 1, f"={g['tot_cre']}").number_format = NUM
                ws.cell(r, k + 2, f"={g['tot_deb']}").number_format = NUM
                k += 3
            r += 1
            ws.cell(r, 1, "Diferencia (debe ser 0)").font = B
            k = 2
            for _ in rep.bancos:
                for c in (k + 1, k + 2):
                    col = get_column_letter(c)
                    cc = ws.cell(r, c, f"=ROUND({col}{r_ctrl - 1}-{col}{r_ctrl},2)")
                    cc.number_format, cc.font = NUM, B
                k += 3
        for c in range(2, ncols + 1):
            ws.column_dimensions[get_column_letter(c)].width = 15
        return r + 3

    r = tabla(4, "Todos los movimientos del banco", solo_banco=False)
    tabla(r, "Movimientos del banco SIN asiento en el libro (estado «Solo en banco»)", solo_banco=True)
    ws.freeze_panes = "B6"


# ------------------------------------------------------------------ Asientos resumen
def _hoja_asientos_resumen(ws, rep: ReporteCliente, rg: dict[str, dict[str, str]]) -> None:
    ws["A1"] = f"{rep.cliente} — Asientos resumen del libro (sin monto en Bs.)"
    ws["A1"].font = TIT
    ws["A2"] = ("Estos asientos registran en US$ el total de cobros del mes y no traen el monto en Bs., así que "
                "no se pueden conciliar movimiento por movimiento. Al lado: las entradas del mismo banco que no "
                "tienen asiento individual, por naturaleza, y la tasa Bs./US$ que resultaría si el asiento "
                "correspondiera a esa naturaleza. Sirve de guía: la decisión es del analista.")
    ws["A2"].font = NOTA
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:H3")
    ws.row_dimensions[2].height = 30
    for col, w in zip("ABCDEFGH", [10, 11, 30, 13, 36, 18, 16, 60]):
        ws.column_dimensions[col].width = w
    _enc(ws, 5, ["Banco", "Fecha", "Asiento", "Monto (US$)", "Naturaleza (entradas sin asiento)",
                 "Monto banco (Bs.)", "Tasa resultante (Bs./US$)", "Nota"])
    r = 6
    for b in rep.bancos:
        rc = rep.conciliaciones[b]
        g = rg[b]
        nat_ent = [n for n in ORDEN if n in COBROS and any(
            x.naturaleza == n and rc.estado_mov[x.indice].estado is EP.SOLO_BANCO for x in rep.clasificados[b])]
        for p in rc.partidas:
            if p.estado is not EP.RESUMEN_SIN_MONTO:
                continue
            a = rc.libro.asientos[p.asientos[0]]
            filas_nat = nat_ent if not a.es_salida else [None]
            for k, n in enumerate(filas_nat):
                vals = [NOMBRE.get(b, b), a.fecha, f"{a.referencia} — {a.descripcion}", a.monto_usd]
                if n is None:
                    vals += ["(asiento de salida)", None, None]
                else:
                    vals += [n, f'=SUMIFS({g["cre"]},{g["nat"]},E{r},{g["est_b"]},"{EP.SOLO_BANCO.value}")',
                             f"=IFERROR(F{r}/D{r},\"\")"]
                vals.append((p.nota or None) if k == 0 else None)
                _fila(ws, r, vals, {2: FECHA, 4: NUM, 6: NUM, 7: TASA}, fill=FILL[EP.RESUMEN_SIN_MONTO] if k == 0 else None)
                r += 1
        ws.cell(r, 1, f"Referencia {NOMBRE.get(b, b)}: tasas implícitas del libro en el mes").font = NOTA
        ws.cell(r, 6, "mín.").font = NOTA
        c = ws.cell(r, 7, f"=MIN({g['tasa']})")
        c.number_format, c.font = TASA, NOTA
        r += 1
        ws.cell(r, 6, "máx.").font = NOTA
        c = ws.cell(r, 7, f"=MAX({g['tasa']})")
        c.number_format, c.font = TASA, NOTA
        r += 2


# ------------------------------------------------------------------ Cierre de caja
def _hoja_caja(ws, rep: ReporteCliente) -> None:
    rc = rep.caja
    ws["A1"] = f"{rep.cliente} — Cierre de caja vs. bancos y libro ({rep.periodo})"
    ws["A1"].font = TIT
    ws["A2"] = (f"Archivo {rc.cierre.archivo}, hoja «{rc.cierre.hoja}». El desglose de cada día sale de las fórmulas "
                "de la celda (=a+b+c). Tarjetas: el lote del día D se abona en D+2; se concilia por el total del "
                "día (crédito + débito) porque el reparto de la caja no coincide con el del banco. Las agrupaciones "
                "se hallan por suma exacta; en días con muchos pagos pequeños la combinación puede no ser única.")
    ws["A2"].font = NOTA
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:J3")
    ws.row_dimensions[2].height = 30
    for col, w in zip("ABCDEFGHIJ", [11, 24, 10, 16, 16, 14, 34, 9, 60, 60]):
        ws.column_dimensions[col].width = w

    r = 4
    for fila, antes, despues in rc.cierre.correcciones:
        ws.cell(r, 1, f"Corrección aplicada (aceptada en la revisión): fila {fila} del cierre, fecha "
                      f"{antes:%d/%m/%Y} → {despues:%d/%m/%Y}. El archivo original no se modificó.").font = ROJA
        r += 1

    # 1. Asientos resumen del libro vs. caja
    r += 1
    ws.cell(r, 1, "1. Asientos resumen del libro frente al cierre de caja (US$)").font = SUB
    r += 1
    _enc(ws, r, ["Banco", "Asiento del libro", "Medio en caja", "US$ libro (asiento o grupo)", "US$ caja (mes)",
                 "Diferencia US$", "Resultado", "Cubre hasta", "Nota"])
    r += 1
    for c in rc.asientos:
        a = rep.conciliaciones[c.banco].libro.asientos[c.asiento]
        _fila(ws, r, [NOMBRE.get(c.banco, c.banco), f"{a.referencia} — {a.descripcion}", c.medio or "-",
                      c.usd_libro_total if c.usd_libro_total is not None else c.usd_libro, c.usd_caja_mes,
                      f"=ROUND(D{r}-E{r},2)",
                      "Conciliado" if c.conciliado else "No coincide", c.hasta, c.nota],
              {4: NUM, 5: NUM, 6: NUM, 8: FECHA},
              fill=FILL[EP.CONCILIADO_CAJA] if c.conciliado and "Faltan" not in c.nota
              else FILL[EP.CAJA_SIN_LIBRO] if c.conciliado else FILL[EP.SOLO_LIBRO])
        r += 1

    # 1b. Tarjetas por total del mes
    for t in rc.totales_mes:
        r += 2
        ws.cell(r, 1, f"Tarjetas por total del mes: {' + '.join(t.medios)} ({', '.join(NOMBRE.get(b, b) for b in t.bancos)})").font = SUB
        filas_t = [("Ventas con tarjeta según caja (mes)", t.caja, False),
                   ("Abonos POS en los bancos durante el período", t.abonado, False),
                   ("(−) Abonos que son ventas del mes anterior", t.mes_anterior, False),
                   ("Abonos sin fecha de venta identificable", t.sin_fecha, False),
                   ("Por abonar después del cierre (caja − abonos del mes)", f"=ROUND(B{r + 1}-(B{r + 2}-B{r + 3}),2)", True)]
        for k, (txt, v, bold) in enumerate(filas_t, 1):
            ws.cell(r + k, 1, txt).font = B if bold else F
            c = ws.cell(r + k, 2, v)
            c.number_format, c.font = NUM, (B if bold else AZUL)
        r += len(filas_t) + 1
        ws.cell(r, 1, "Verificar con el estado de cuenta del mes siguiente que lo «por abonar» llegue completo.").font = NOTA

    # 2. Resumen por medio
    r += 2
    ws.cell(r, 1, "2. Cierre de caja frente al banco: resumen por medio (Bs.)").font = SUB
    r += 1
    estados_cols = [("Conciliado", sorted(CONCILIADOS_CAJA, key=lambda e: e.value)),
                    ("Posible error", [EstadoCaja.POSIBLE_ERROR]),
                    ("No encontrado", [EstadoCaja.NO_ENCONTRADO]),
                    ("En tránsito", [EstadoCaja.EN_TRANSITO]),
                    ("Diferencia en el día", [EstadoCaja.DIFERENCIA_DIA]),
                    ("Por total del mes", [EstadoCaja.TOTAL_MES])]
    _enc(ws, r, ["Medio", "Total caja"] + [t for t, _ in estados_cols] + ["Control (debe ser 0)"])
    r += 1
    r_resumen = r
    medios = sorted({l.medio for l in rc.lineas})
    r_det_ini = r_resumen + len(medios) + 5
    r_det_fin = r_det_ini + len(rc.lineas) - 1
    rng = lambda col: f"${col}${r_det_ini}:${col}${r_det_fin}"
    for m in medios:
        ws.cell(r, 1, m).font = F
        ws.cell(r, 2, f'=SUMIFS({rng("D")},{rng("B")},$A{r})')
        for k, (_, ests) in enumerate(estados_cols, 3):
            ws.cell(r, k, "=" + "+".join(f'SUMIFS({rng("D")},{rng("B")},$A{r},{rng("G")},"{e.value}")' for e in ests))
        ws.cell(r, 9, f"=ROUND(B{r}-SUM(C{r}:H{r}),2)")
        for k in range(1, 10):
            ws.cell(r, k).border = LINEA
            if k > 1:
                ws.cell(r, k).number_format = NUM
                ws.cell(r, k).font = F
        r += 1

    # 3. Detalle
    r = r_det_ini - 2
    ws.cell(r, 1, "3. Detalle día por día (pagos móviles por sumando; tarjetas por día)").font = SUB
    r += 1
    _enc(ws, r, ["Fecha caja", "Medio", "Banco", "Monto caja (Bs.)", "Monto banco (Bs.)", "Diferencia (Bs.)",
                 "Estado", "¿En libro?", "Movimientos del banco", "Nota"])
    r += 1
    assert r == r_det_ini
    for l in rc.lineas:
        movs = [rep.conciliaciones[b].extracto.movimientos[i] for b, i in l.movs]
        monto_b = sum((m.credito or m.debito for m in movs), Decimal("0.00")) if movs else None
        _fila(ws, r, [l.fecha, l.medio, NOMBRE.get(l.banco, l.banco or "-"), l.monto, monto_b,
                      f'=IF(E{r}="","",ROUND(E{r}-D{r},2))', l.estado.value,
                      "Sí" if l.en_libro else "No",
                      " + ".join(f"{m.fecha_contable:%d/%m} {m.descripcion} {m.credito or m.debito}" for m in movs)[:500]
                      or None, l.nota or None],
              {1: FECHA, 4: NUM, 5: NUM, 6: NUM}, fill=FILL_CAJA.get(l.estado))
        r += 1
    ws.cell(r, 2, "TOTAL").font = B
    for col in "DE":
        c = ws[f"{col}{r}"]
        c.value = f"=SUM({col}{r_det_ini}:{col}{r - 1})"
        c.font, c.number_format, c.fill = B, NUM, TOT

    # 4. Banco sin caja
    r += 3
    ws.cell(r, 1, "4. Créditos del banco sin registro en el cierre de caja").font = SUB
    r += 1
    _enc(ws, r, ["Fecha banco", "Descripción", "Banco", "Monto (Bs.)", "", "", "Motivo"])
    r += 1
    r0 = r
    for b, i, motivo in rc.sin_caja:
        m = rep.conciliaciones[b].extracto.movimientos[i]
        _fila(ws, r, [m.fecha_contable, f"{m.descripcion} {m.referencia}".strip(), NOMBRE.get(b, b),
                      m.credito, None, None, motivo], {1: FECHA, 4: NUM})
        r += 1
    if r > r0:
        ws.cell(r, 2, "TOTAL").font = B
        c = ws.cell(r, 4, f"=SUM(D{r0}:D{r - 1})")
        c.font, c.number_format, c.fill = B, NUM, TOT


# ------------------------------------------------------------------ Revisión (hoja editable)
def _hoja_revision(ws, rep: ReporteCliente, rv) -> None:
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
    from .revision import ABIERTAS, DECISIONES, Situacion

    abiertas = rv.abiertas
    ws["A1"] = f"{rep.cliente} — Revisión de la conciliación {rep.periodo}"
    ws["A1"].font = TIT
    ws["A2"] = (f"Corrida {rv.corrida} · generada {rv.generado:%d/%m/%Y %H:%M} · "
                f"{len(rv.pendientes)} partidas, {len(abiertas)} abiertas")
    ws["A2"].font = NOTA
    if rv.ok_general:
        ws["A3"] = "OK GENERAL: todos los bancos cuadran y no quedan partidas abiertas. El analista puede cerrar el período."
        ws["A3"].font = Font(name="Arial", size=12, bold=True, color="008000")
    else:
        motivo = []
        if rv.bancos_no_cuadran:
            motivo.append("no cuadra: " + ", ".join(NOMBRE.get(b, b) for b in rv.bancos_no_cuadran))
        if abiertas:
            motivo.append(f"{len(abiertas)} partidas abiertas")
        ws["A3"] = "CON PENDIENTES — " + "; ".join(motivo)
        ws["A3"].font = Font(name="Arial", size=12, bold=True, color="C00000")
    ws["A4"] = ("Cómo usar: lea «Explicación» y «Qué hacer» de cada partida. Complete solo las columnas en amarillo "
                "(Decisión, Comentario, Revisado por; «Decisión sugerida» es una guía) y vuelva a "
                "subir ESTE archivo junto con los exports corregidos del sistema, si los hay. Decisiones: "
                "«Aceptar» (lo propuesto es correcto), «Justificado» (correcto así; comentario obligatorio), "
                "«Corregido en el sistema» (se corrigió el asiento; en la próxima corrida debe desaparecer). "
                "No cambie la columna Código.")
    ws["A4"].font = NOTA
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A4:J5")
    ws.row_dimensions[4].height = 32
    ws.row_dimensions[5].height = 32
    for a in rv.advertencias:
        ws["A6"] = a
        ws["A6"].font = ROJA

    # Conteo por situación
    r = 7
    from collections import Counter
    cnt = Counter(p.situacion for p in rv.pendientes)
    ws.cell(r, 4, "Situación").font = B
    ws.cell(r, 5, "Partidas").font = B
    for k, sit in enumerate(Situacion, r + 1):
        ws.cell(k, 4, sit.value).font = F
        ws.cell(k, 5, cnt.get(sit, 0)).font = ROJA if sit in ABIERTAS and cnt.get(sit, 0) else F
    k += 1
    ws.cell(k, 4, "Resueltas desde la corrida anterior").font = F
    ws.cell(k, 5, len(rv.resueltas)).font = F

    r = k + 2
    titulos = ["Código", "Banco", "Origen", "Tipo", "Fecha", "Descripción", "Monto (Bs.)", "Monto (US$)",
               "Explicación", "Qué hacer", "Decisión sugerida", "Situación", "Decisión", "Comentario",
               "Revisado por", "Aviso", "Detalle técnico"]
    _enc(ws, r, titulos, [15, 9, 8, 30, 11, 36, 15, 12, 70, 45, 20, 20, 22, 40, 16, 40, 50])
    editable = PatternFill("solid", start_color="FFF2CC")
    fill_sit = {Situacion.CERRADO: PatternFill("solid", start_color="E2EFDA"),
                Situacion.CORREGIDO_SIGUE: PatternFill("solid", start_color="FFC7CE"),
                Situacion.INCOMPLETA: PatternFill("solid", start_color="FFEB9C")}
    orden_sit = {Situacion.CORREGIDO_SIGUE: 0, Situacion.INCOMPLETA: 1, Situacion.NUEVO: 2,
                 Situacion.SIN_DECISION: 3, Situacion.CERRADO: 4}
    ajuste = Alignment(wrap_text=True, vertical="top")
    r0 = r + 1
    for p in sorted(rv.pendientes, key=lambda p: (orden_sit[p.situacion], p.banco, p.tipo, p.fecha or dt.date.min)):
        r += 1
        _fila(ws, r, [p.codigo, NOMBRE.get(p.banco, p.banco), p.origen, p.tipo, p.fecha, p.descripcion,
                      p.monto_bs, p.monto_usd, (p.explicacion or "")[:1500] or None, p.que_hacer or None,
                      p.sugerencia, p.situacion.value, p.decision or None, p.comentario or None,
                      p.revisado_por or None, p.aviso or None, (p.detalle or "")[:800] or None],
              {5: FECHA, 7: NUM, 8: NUM})
        for c in (9, 10):
            ws.cell(r, c).alignment = ajuste
        ws.cell(r, 12).fill = fill_sit.get(p.situacion, PatternFill())
        for c in (13, 14, 15):
            ws.cell(r, c).fill = editable
            ws.cell(r, c).protection = Protection(locked=False)
        if p.aviso:
            ws.cell(r, 16).font = ROJA
    if r >= r0:
        dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONES) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Decisión no válida",
                            error="Elija Aceptar, Justificado o Corregido en el sistema.")
        ws.add_data_validation(dv)
        dv.add(f"M{r0}:M{r}")
        ws.auto_filter.ref = f"A{r0 - 1}:Q{r}"
    ws.cell(r0 - 1, 13).comment = Comment("Elija de la lista. «Justificado» requiere comentario.", "CVP")
    ws.freeze_panes = f"B{r0}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False

    if rv.resueltas:
        r += 3
        ws.cell(r, 1, "Resueltas: estaban en la corrida anterior y ya no aparecen").font = SUB
        r += 1
        _enc(ws, r, ["Código", "", "", "Tipo", "", "Descripción", "", "", "", "", "", "", "Decisión anterior",
                     "Comentario", "Revisado por"])
        for d in rv.resueltas:
            r += 1
            _fila(ws, r, [d.codigo, None, None, d.tipo, None, d.descripcion, None, None, None, None, None, None,
                          d.decision or "—", d.comentario or None, d.revisado_por or None],
                  fill=PatternFill("solid", start_color="E2EFDA"))


def _hoja_historial(ws, rep: ReporteCliente, rv) -> None:
    ws["A1"] = f"{rep.cliente} — Historial de decisiones {rep.periodo}"
    ws["A1"].font = TIT
    _enc(ws, 3, ["Corrida", "Procesado", "Código", "Tipo", "Descripción", "Decisión", "Comentario",
                 "Revisado por", "Resultado"], [8, 17, 15, 34, 40, 22, 40, 16, 50])
    r = 3
    for fila in rv.historial:
        r += 1
        _fila(ws, r, list(fila[:9]), {2: "DD/MM/YYYY HH:MM"})
    if r == 3:
        ws.cell(4, 1, "Primera corrida: todavía no hay decisiones.").font = NOTA
    ws.protection.sheet = True


def _hoja_control(ws, rep: ReporteCliente, rv) -> None:
    from .revision import VERSION
    filas = [("version", VERSION), ("cliente", rep.cliente), ("periodo", rep.periodo),
             ("corrida", rv.corrida), ("generado", rv.generado.isoformat()),
             ("ok_general", "SI" if rv.ok_general else "NO"), ("codigos_desde_fila", 10)]
    for i, (k, v) in enumerate(filas, 1):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
    i = 10
    for p in rv.pendientes:
        ws.cell(i, 1, "codigo")
        ws.cell(i, 2, p.codigo)
        i += 1
    for p in rv.pendientes:
        if p.propuesta:
            ws.cell(i, 1, "propuesta")
            ws.cell(i, 2, p.codigo)
            ws.cell(i, 3, p.propuesta[0])
            ws.cell(i, 4, p.propuesta[1].isoformat())
            i += 1
    if rep.caja:
        for fila, _, nueva in rep.caja.cierre.correcciones:
            ws.cell(i, 1, "correccion")
            ws.cell(i, 2, fila)
            ws.cell(i, 3, nueva.isoformat())
            i += 1
    ws.sheet_state = "hidden"
    ws.protection.sheet = True


# ------------------------------------------------------------------ Divisas
def _hoja_divisas(ws, rep: ReporteCliente) -> None:
    ws["A1"] = f"{rep.cliente} — Cuentas en divisas ({rep.periodo})"
    ws["A1"].font = TIT
    kx = rep.kardex
    ws["A2"] = ("Ventas (cierre de caja) ↔ libro del sistema ↔ Kardex"
                + (f" ({kx.archivo}, hoja {kx.hoja})" if kx else "") + ". Todos los montos en US$.")
    ws["A2"].font = NOTA
    for col, w in zip("ABCDEFGHI", [26, 12, 46, 14, 14, 14, 30, 60, 14]):
        ws.column_dimensions[col].width = w
    r = 4
    _enc(ws, r, ["Cuenta", "", "Control", "Libro", "Ventas / Kardex", "Diferencia", "Resultado", "Nota"])
    for res in rep.divisas:
        n = res.cuenta.nombre
        filas = []
        if res.ventas_mes is not None:
            a = res.libro.asientos[res.asiento_ventas] if res.asiento_ventas is not None else None
            filas.append(("Ventas del mes = asiento de ventas", a.debito_usd if a else None, res.ventas_mes,
                          res.ventas_ok, res.ventas_nota))
        for t, (lib, k) in (("Saldo inicial libro = VIENEN del Kardex", res.saldo_ini),
                            ("Saldo final libro = DISPONIBLE del Kardex", res.saldo_fin)):
            if k is not None:
                filas.append((t, lib, k, lib == k, ""))
        if res.dias:
            mal = res.dias_con_diferencia
            filas.append((f"Ventas día a día = Kardex ({len(res.dias)} días)", None, None, not mal,
                          "; ".join(f"{d.fecha:%d/%m} dif. {d.diferencia}" for d in mal)
                          or "; ".join(d.nota for d in res.dias if d.nota)))
        from collections import Counter
        cnt = Counter(p.estado for p in res.partidas)
        filas.append((f"Movimientos del libro ↔ Kardex", None, None,
                      not (cnt.get("Solo en libro") or cnt.get("Solo en Kardex")),
                      ", ".join(f"{k}: {v}" for k, v in cnt.items())))
        for t, lib, k, ok, nota in filas:
            r += 1
            _fila(ws, r, [n, None, t, lib, k, f'=IF(OR(D{r}="",E{r}=""),"",ROUND(D{r}-E{r},2))',
                          "OK" if ok else "REVISAR", nota or None], {4: NUM, 5: NUM, 6: NUM},
                  fill=FILL[EP.CONCILIADO] if ok else FILL[EP.DIFERENCIA_MONTO])
    if rep.cruces_divisas_bancos:
        r += 2
        ws.cell(r, 1, "Venta de divisas ↔ bancos").font = SUB
        for t in rep.cruces_divisas_bancos:
            r += 1
            ws.cell(r, 1, t).font = ROJA if "diferencia US$ 0.00" not in t else F

    # Detalle de movimientos
    r += 3
    ws.cell(r, 1, "Detalle: movimientos del libro y su pareja en el Kardex").font = SUB
    r += 1
    _enc(ws, r, ["Cuenta", "Fecha", "Descripción (libro)", "Entrada US$", "Salida US$", "Kardex US$",
                 "Estado", "Kardex / nota"])
    for res in rep.divisas:
        for p in sorted(res.partidas, key=lambda p: (res.libro.asientos[p.asientos[0]].fecha if p.asientos
                                                    else p.kardex[0].fecha)):
            r += 1
            a = res.libro.asientos[p.asientos[0]] if p.asientos else None
            desc = " + ".join(res.libro.asientos[j].descripcion for j in p.asientos) or "—"
            _fila(ws, r, [res.cuenta.nombre, a.fecha if a else p.kardex[0].fecha, desc,
                          sum((res.libro.asientos[j].debito_usd for j in p.asientos), Decimal("0.00")) or None,
                          sum((res.libro.asientos[j].credito_usd for j in p.asientos), Decimal("0.00")) or None,
                          sum((k.monto for k in p.kardex), Decimal("0.00")) if p.kardex else None, p.estado,
                          p.nota or " + ".join(f"fila {k.fila} {k.descripcion}" for k in p.kardex) or None],
                  {2: FECHA, 4: NUM, 5: NUM, 6: NUM},
                  fill=FILL[EP.CONCILIADO] if p.estado.startswith("Conciliado") else FILL[EP.DIFERENCIA_MONTO])
    ws.freeze_panes = "A5"

"""Servicio «Solo conversión»: estados de cuenta (PDF / Excel del banco) → Excel, con resumen por ítem.

Sin libro del sistema ni conciliación. Por cada cuenta:
  1. Se lee el estado de cuenta con el lector del banco.
  2. Se cuadra contra lo que informa el propio banco (saldo anterior, totales, saldo final, saldo línea a
     línea). Si no cuadra se dice exactamente dónde; nada se ajusta.
  3. Cada movimiento se clasifica por naturaleza (cobros POS, pago móvil, transferencias, comisiones…)
     y por concepto (la descripción del banco sin números de referencia).

Excel:
  Resumen               Cuadre de cada cuenta (saldo inicial, débitos, créditos, saldo final) y avisos.
  Resumen por ítem      Por cuenta: naturaleza → concepto, con nº de movimientos, débitos, créditos y neto;
                        subtotales y un control que debe dar 0 contra la hoja de movimientos.
                        Al final, el consolidado de todas las cuentas por naturaleza.
  Mov <cuenta>          Todos los movimientos con naturaleza, concepto y saldo calculado (fórmulas).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

from .cargos import tipo_de
from .cuadre import Estado, ResultadoCuadre, cuadrar_extracto
from .lectores import LECTORES
from .modelo import Extracto
from .montos import CERO
from .naturaleza import ORDEN, MovClasificado, clasificar_extracto, marcar_traslados


@dataclass
class CuentaConvertida:
    nombre: str                       # nombre de la cuenta en la app (p. ej. «BNC ***1800»)
    banco: str
    extracto: Extracto
    cuadre: ResultadoCuadre
    movs: list[MovClasificado]


@dataclass
class Item:
    naturaleza: str
    concepto: str
    n: int = 0
    debitos: Decimal = CERO
    creditos: Decimal = CERO

    @property
    def neto(self) -> Decimal:
        return self.creditos - self.debitos


@dataclass
class Conversion:
    cliente: str
    periodo: str
    cuentas: list[CuentaConvertida] = field(default_factory=list)

    @property
    def todo_cuadra(self) -> bool:
        return all(c.cuadre.estado is not Estado.NO_CUADRA for c in self.cuentas)


def convertir(cliente: str, periodo: str, archivos: list[tuple[str, str, Path]]) -> Conversion:
    """archivos: [(nombre de la cuenta, banco, estado de cuenta), ...]. Admite varias cuentas del mismo banco."""
    conv = Conversion(cliente, periodo)
    for nombre, banco, ruta in archivos:
        if banco not in LECTORES:
            raise ValueError(f"«{nombre}»: el banco {banco} todavía no tiene lector (disponibles: {', '.join(LECTORES)}).")
        res = LECTORES[banco].leer([Path(ruta)])
        if len(res.extractos) != 1:
            raise ValueError(f"«{nombre}»: se esperaba un estado de cuenta en {Path(ruta).name}, "
                             f"se encontraron {len(res.extractos)}.")
        e = res.extractos[0]
        conv.cuentas.append(CuentaConvertida(nombre, banco, e, cuadrar_extracto(e), clasificar_extracto(e)))
    marcar_traslados([c.movs for c in conv.cuentas])
    return conv


def items(movs: list[MovClasificado]) -> list[Item]:
    """Resumen por naturaleza (en el orden habitual) y concepto (de mayor a menor monto)."""
    por: dict[tuple[str, str], Item] = {}
    for x in movs:
        it = por.setdefault((x.naturaleza, tipo_de(x.mov)), Item(x.naturaleza, tipo_de(x.mov)))
        it.n += 1
        it.debitos += x.mov.debito
        it.creditos += x.mov.credito
    orden = {n: i for i, n in enumerate(ORDEN)}
    return sorted(por.values(), key=lambda it: (orden.get(it.naturaleza, 99), it.naturaleza,
                                                -(it.debitos + it.creditos), it.concepto))


# ---------------------------------------------------------------- Excel
def _hoja_nombre(nombre: str, usados: set[str]) -> str:
    base = "Mov " + re.sub(r"[\[\]:*?/\\']", " ", nombre).strip()
    base = base[:31]
    n, k = base, 2
    while n in usados:
        suf = f" ({k})"
        n, k = base[:31 - len(suf)] + suf, k + 1
    usados.add(n)
    return n


def excel_conversion(conv: Conversion, salida: Path) -> Path:
    from .exportar import B, BLANCA, ENC, F, FECHA, LINEA, NOTA, NUM, ROJA, SUB, TIT, TOT, VERDE_TXT, _enc, _fila, _q

    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_it = wb.create_sheet("Resumen por ítem")
    usados: set[str] = {"Resumen", "Resumen por ítem"}
    ref: dict[int, dict[str, str]] = {}

    # ---- Movimientos por cuenta
    for k, c in enumerate(conv.cuentas):
        nombre = _hoja_nombre(c.nombre, usados)
        ws = wb.create_sheet(nombre)
        e = c.extracto
        ws["A1"] = f"{e.titular or conv.cliente} — {c.nombre} ({e.cuenta})"
        ws["A1"].font = TIT
        ws["A2"] = (f"Período {e.desde:%d/%m/%Y} al {e.hasta:%d/%m/%Y} · " if e.desde else f"Al {e.hasta:%d/%m/%Y} · ") \
            + f"Archivo {e.archivo} · Saldo anterior:"
        ws["A2"].font = NOTA
        ws["H2"] = e.saldo_anterior
        ws["H2"].number_format = NUM
        _enc(ws, 4, ["Fecha", "Fecha valor", "Referencia", "Descripción", "Detalle", "Débito (Bs.)", "Crédito (Bs.)",
                     "Saldo banco (Bs.)", "Saldo calculado (Bs.)", "Naturaleza", "Concepto", "Nota"],
             [11, 11, 16, 34, 36, 15, 15, 16, 16, 30, 32, 40])
        r = 5
        for x in c.movs:
            m = x.mov
            prev = "$H$2" if r == 5 else f"I{r - 1}"
            _fila(ws, r, [m.fecha_contable, m.fecha_real if m.fecha_real and m.fecha_real != m.fecha_contable else None,
                          m.referencia or None, m.descripcion, m.detalle or None, m.debito or None, m.credito or None,
                          m.saldo, f"=ROUND({prev}-F{r}+G{r},2)", x.naturaleza, tipo_de(m), x.contraparte or None],
                  {1: FECHA, 2: FECHA, 6: NUM, 7: NUM, 8: NUM, 9: NUM})
            r += 1
        ultima = r - 1
        ws.cell(r, 4, "TOTALES").font = B
        for col in "FG":
            cel = ws[f"{col}{r}"]
            cel.value = f"=SUM({col}5:{col}{ultima})" if ultima >= 5 else 0
            cel.font, cel.number_format, cel.fill = B, NUM, TOT
        ws.cell(r + 1, 4, "Control: saldo calculado − saldo del banco en la última línea (debe ser 0)").font = NOTA
        if ultima >= 5:
            cel = ws.cell(r + 1, 9, f"=ROUND(I{ultima}-H{ultima},2)")
            cel.font, cel.number_format = B, NUM
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:L{max(ultima, 4)}"
        h = _q(nombre)
        ref[k] = {"deb": f"{h}!F{r}", "cre": f"{h}!G{r}", "hoja": nombre}

    # ---- Resumen (cuadre)
    ws = ws_res
    ws["A1"] = f"{conv.cliente} — Estados de cuenta {conv.periodo}"
    ws["A1"].font = TIT
    ws["A2"] = ("Conversión de los estados de cuenta a Excel. «Cuadra» = el detalle reproduce el saldo anterior, los "
                "totales y el saldo final que informa el banco. Nada se ajusta.")
    ws["A2"].font = NOTA
    _enc(ws, 4, ["Cuenta", "Banco", "Nº cuenta", "Movimientos", "Saldo anterior", "Débitos", "Créditos",
                 "Saldo final (calculado)", "Saldo final (banco)", "Estado"],
         [30, 14, 24, 13, 17, 17, 17, 19, 19, 22])
    r = 4
    for k, c in enumerate(conv.cuentas):
        r += 1
        e, q = c.extracto, c.cuadre
        fin_banco = e.totales_banco.nuevo_saldo if e.totales_banco and e.totales_banco.nuevo_saldo is not None \
            else (e.movimientos[-1].saldo if e.movimientos and e.movimientos[-1].saldo is not None else None)
        _fila(ws, r, [c.nombre, c.banco, e.cuenta, len(e.movimientos), e.saldo_anterior, f"={ref[k]['deb']}",
                      f"={ref[k]['cre']}", f"=ROUND(E{r}-F{r}+G{r},2)", fin_banco,
                      {"cuadra": "Cuadra", "requiere_revision": "Cuadra (revisar avisos)",
                       "no_cuadra": "NO CUADRA"}[q.estado.value]],
              {5: NUM, 6: NUM, 7: NUM, 8: NUM, 9: NUM})
        ws.cell(r, 10).font = ROJA if q.estado is Estado.NO_CUADRA else VERDE_TXT
    r += 2
    hay = False
    for c in conv.cuentas:
        msgs = [f"✗ {d.mensaje}" for d in c.cuadre.diferencias] + [f"• {a.mensaje}" for a in c.cuadre.avisos] \
            + [f"✗ Lectura: {x}" for x in c.extracto.errores_lectura] + [f"• {x}" for x in c.extracto.avisos_lectura]
        if not msgs:
            continue
        if not hay:
            ws.cell(r, 1, "Diferencias y avisos").font = SUB
            r += 1
            hay = True
        ws.cell(r, 1, c.nombre).font = B
        for msg in msgs:
            r += 1
            ws.cell(r, 2, msg).font = ROJA if msg.startswith("✗") else F
        r += 1

    # ---- Resumen por ítem
    ws = ws_it
    ws["A1"] = f"{conv.cliente} — Resumen por ítem {conv.periodo}"
    ws["A1"].font = TIT
    ws["A2"] = ("Movimientos agrupados por naturaleza y concepto (descripción del banco sin números de referencia). "
                "El control de cada cuenta compara contra la hoja de movimientos y debe dar 0.")
    ws["A2"].font = NOTA
    for col, w in zip("ABCDEFG", [34, 40, 12, 18, 18, 18, 30]):
        ws.column_dimensions[col].width = w
    r = 4
    consolidado: dict[str, Item] = {}
    for k, c in enumerate(conv.cuentas):
        ws.cell(r, 1, c.nombre).font = SUB
        r += 1
        for i, t in enumerate(["Naturaleza", "Concepto", "Movimientos", "Débitos (Bs.)", "Créditos (Bs.)", "Neto (Bs.)"], 1):
            cel = ws.cell(r, i, t)
            cel.font, cel.fill = BLANCA, ENC
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        subt: list[int] = []
        lista = items(c.movs)
        idx = 0
        while idx < len(lista):
            nat = lista[idx].naturaleza
            r0 = r + 1
            while idx < len(lista) and lista[idx].naturaleza == nat:
                it = lista[idx]
                r += 1
                _fila(ws, r, [nat if r == r0 else None, it.concepto, it.n, it.debitos or None, it.creditos or None,
                              f"=ROUND(N(E{r})-N(D{r}),2)"], {4: NUM, 5: NUM, 6: NUM})
                cons = consolidado.setdefault(nat, Item(nat, ""))
                cons.n += it.n
                cons.debitos += it.debitos
                cons.creditos += it.creditos
                idx += 1
            r += 1
            subt.append(r)
            _fila(ws, r, [f"Subtotal {nat}", None, f"=SUM(C{r0}:C{r - 1})", f"=SUM(D{r0}:D{r - 1})",
                          f"=SUM(E{r0}:E{r - 1})", f"=ROUND(E{r}-D{r},2)"], {4: NUM, 5: NUM, 6: NUM}, fill=TOT, font=B)
        r += 1
        suma = lambda col: "=" + ("+".join(f"{col}{x}" for x in subt) if subt else "0")
        _fila(ws, r, [f"TOTAL {c.nombre}", None, suma("C"), suma("D"), suma("E"), f"=ROUND(E{r}-D{r},2)"],
              {4: NUM, 5: NUM, 6: NUM}, font=B)
        r += 1
        _fila(ws, r, ["Control contra la hoja de movimientos (debe ser 0)", None, None,
                      f"=ROUND(D{r - 1}-{ref[k]['deb']},2)", f"=ROUND(E{r - 1}-{ref[k]['cre']},2)", None],
              {4: NUM, 5: NUM}, font=NOTA)
        r += 3
    if len(conv.cuentas) > 1:
        ws.cell(r, 1, "Todas las cuentas").font = SUB
        r += 1
        for i, t in enumerate(["Naturaleza", "", "Movimientos", "Débitos (Bs.)", "Créditos (Bs.)", "Neto (Bs.)"], 1):
            cel = ws.cell(r, i, t)
            cel.font, cel.fill = BLANCA, ENC
        r0 = r + 1
        orden = {n: i for i, n in enumerate(ORDEN)}
        for nat in sorted(consolidado, key=lambda n: (orden.get(n, 99), n)):
            it = consolidado[nat]
            r += 1
            _fila(ws, r, [nat, None, it.n, it.debitos or None, it.creditos or None, f"=ROUND(N(E{r})-N(D{r}),2)"],
                  {4: NUM, 5: NUM, 6: NUM})
        r += 1
        _fila(ws, r, ["TOTAL", None, f"=SUM(C{r0}:C{r - 1})", f"=SUM(D{r0}:D{r - 1})", f"=SUM(E{r0}:E{r - 1})",
                      f"=ROUND(E{r}-D{r},2)"], {4: NUM, 5: NUM, 6: NUM}, fill=TOT, font=B)
    ws.freeze_panes = "A4"

    salida = Path(salida)
    wb.save(salida)
    return salida

"""Lector del cierre de caja diario del restaurante (Excel "VENTAS ...").

Formato observado (WEI REST, agosto 2026, hoja "VENTAS 2024"):
  - Fila 3: encabezados de grupo por medio de pago ("PAGO MOVIL BNC", "T. CREDITO BNC", ...).
  - Fila 4: subencabezados ("LOTE", "BOLIVARES", "DOLARES").
  - Desde la fila 5: un día por fila. Columna B = fecha, D = tasa del día.
  - Cuando un día tuvo varios pagos, la celda en Bs. es una fórmula con el desglose:
    =37978.04+70926.78+232751.12+81000. Esos sumandos son los pagos que aparecen
    uno por uno en el estado de cuenta.
  - Fila "Total Mensual": totales del mes. Los US$ de cada día son Bs. / tasa.

Se leen el archivo dos veces: con fórmulas (para el desglose) y con valores (para los totales).
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Optional

import openpyxl

from .montos import CERO, MontoInvalido, monto_excel

# Encabezado de grupo → (banco, tipo). El tipo decide cómo se concilia.
BANCOS_POR_NOMBRE = {"BNC": "BNC", "100% BANCO": "100_BANCO", "BANPLUS": "BANPLUS"}
POS, PAGO_MOVIL, VUELTO = "pos", "pago_movil", "vuelto"


def _medio(encabezado: str) -> Optional[tuple[Optional[str], str, str]]:
    """'PAGO MOVIL BNC' -> ('BNC', 'pago_movil', 'Pago móvil BNC'); None si no interesa."""
    t = re.sub(r"\s+", " ", encabezado.strip().upper())
    m = re.match(r"PAGO MOVIL (.+)$", t)
    if m and m.group(1) in BANCOS_POR_NOMBRE:
        return BANCOS_POR_NOMBRE[m.group(1)], PAGO_MOVIL, f"Pago móvil {m.group(1)}"
    m = re.match(r"T\. ?(CREDITO|DEBITO) (.+)$", t)
    if m and m.group(2) in BANCOS_POR_NOMBRE:
        return BANCOS_POR_NOMBRE[m.group(2)], POS, f"T. {m.group(1).lower()} {m.group(2)}"
    if t == "VUELTOS":
        return None, VUELTO, "Vueltos"
    return None


@dataclass
class ItemCaja:
    """El monto de un medio de pago en un día, con su desglose."""
    fecha: dt.date
    fila: int
    medio: str               # "Pago móvil BANPLUS", "T. crédito BNC", ...
    banco: Optional[str]
    tipo: str                # pos / pago_movil / vuelto
    monto_bs: Decimal
    usd: Decimal             # Bs. / tasa, tal como lo calcula la hoja (sin redondear a 2 decimales)
    componentes: list[Decimal] = field(default_factory=list)   # [monto_bs] si no hay fórmula
    formula: str = ""
    modo: str = ""           # desglose / lote / total_diario / total_mes (ver clientes.py)
    bancos: list[str] = field(default_factory=list)


@dataclass
class CierreCaja:
    archivo: str
    hoja: str
    tasas: dict[dt.date, Decimal]
    items: list[ItemCaja]
    errores: list[str] = field(default_factory=list)
    # Correcciones de fecha aplicadas (aceptadas por el analista): (fila, fecha en el archivo, fecha usada)
    correcciones: list[tuple[int, dt.date, dt.date]] = field(default_factory=list)
    fechas_faltantes: list[dt.date] = field(default_factory=list)

    def items_de(self, medio: str) -> list[ItemCaja]:
        return [i for i in self.items if i.medio == medio]

    @property
    def medios(self) -> list[str]:
        vistos: list[str] = []
        for i in self.items:
            if i.medio not in vistos:
                vistos.append(i.medio)
        return vistos


_SUMANDOS = re.compile(r"^=\+?\s*(\d+(?:\.\d+)?)(\s*\+\s*\d+(?:\.\d+)?)*\s*$")


def _desglose(formula: object) -> Optional[list[Decimal]]:
    """'=37978.04+70926.78' -> [37978.04, 70926.78]. None si no es una suma de números."""
    if not isinstance(formula, str) or not _SUMANDOS.match(formula.replace(" ", "")):
        return None
    return [monto_excel(x) for x in re.findall(r"\d+(?:\.\d+)?", formula)]


MODO_AUTO = {PAGO_MOVIL: "desglose", POS: "lote", VUELTO: "desglose"}


def _norm(t: object) -> str:
    return re.sub(r"\s+", " ", str(t or "")).strip().upper()


def leer_cierre_caja(ruta: Path, hoja: Optional[str] = None, medios: Optional[list] = None,
                     correcciones: Optional[dict[int, dt.date]] = None) -> CierreCaja:
    """medios: lista de clientes.MedioCaja. Si no se da, las columnas se detectan por el nombre
    («PAGO MOVIL <banco>», «T. CREDITO <banco>», «VUELTOS»).
    correcciones: {fila: fecha correcta} aceptadas por el analista en la hoja Revisión. Se aplican al
    leer; el archivo original no se modifica."""
    correcciones = correcciones or {}
    aplicadas: list[tuple[int, dt.date, dt.date]] = []
    ruta = Path(ruta)
    wv = openpyxl.load_workbook(str(ruta), data_only=True)
    wf = openpyxl.load_workbook(str(ruta))
    if hoja is None:
        hoja = next((n for n in wv.sheetnames if n.upper().startswith("VENTAS")
                     and wv[n].sheet_state == "visible"), None)
        if hoja is None:
            raise ValueError(f"{ruta.name}: no hay una hoja visible que empiece por 'VENTAS'")
    sv, sf = wv[hoja], wf[hoja]

    # Encabezados: fila con "PAGO MOVIL ..." y la siguiente con BOLIVARES/DOLARES.
    buscados = {_norm(m.columna) for m in medios} if medios else set()
    f_grupo = next(r for r in range(1, 15)
                   if any(isinstance(c.value, str) and (_norm(c.value).startswith("PAGO MOVIL")
                                                        or _norm(c.value) in buscados) for c in sv[r]))
    f_sub = f_grupo + 1
    columnas = []  # (col_bs, col_usd, banco, tipo, medio, modo, bancos)
    por_columna = {_norm(m.columna): m for m in medios} if medios else {}
    for c in sv[f_grupo]:
        if not isinstance(c.value, str):
            continue
        if medios:
            m = por_columna.get(_norm(c.value))
            if m is None:
                continue
            med = (m.bancos[0] if len(m.bancos) == 1 else None, m.tipo, m.nombre or c.value.strip(),
                   m.modo, list(m.bancos))
        else:
            med = _medio(c.value)
            if med is None:
                continue
            med = (*med, MODO_AUTO[med[1]], [med[0]] if med[0] else [])
        col = c.column
        while str(sv.cell(f_sub, col).value or "").strip().upper() != "BOLIVARES":
            col += 1
            if col > c.column + 3:
                raise ValueError(f"{ruta.name}: no se encontró la columna BOLIVARES de '{c.value}'")
        columnas.append((col, col + 1, *med))
    c_fecha = next(c.column for c in sv[f_sub] if str(c.value or "").strip() == "Fecha")
    c_tasa = next(c.column for c in sv[f_sub] if str(c.value or "").strip().upper().startswith("TASA"))

    tasas: dict[dt.date, Decimal] = {}
    items: list[ItemCaja] = []
    errores: list[str] = []
    for r in range(f_sub + 1, sv.max_row + 1):
        fecha = sv.cell(r, c_fecha).value
        if not isinstance(fecha, dt.datetime):
            continue
        fecha = fecha.date()
        if r in correcciones and correcciones[r] != fecha:
            aplicadas.append((r, fecha, correcciones[r]))
            fecha = correcciones[r]
        tasa = sv.cell(r, c_tasa).value
        tasas[fecha] = Decimal(repr(tasa)) if isinstance(tasa, (int, float)) else CERO
        for col_bs, col_usd, banco, tipo, medio, modo, bancos in columnas:
            v = sv.cell(r, col_bs).value
            if v in (None, "", 0):
                continue
            try:
                monto = monto_excel(v)
            except MontoInvalido as e:
                errores.append(f"fila {r} {medio}: {e}")
                continue
            if monto == 0:
                continue
            formula = sf.cell(r, col_bs).value
            comps = _desglose(formula)
            if comps is None:
                comps = [monto]
            elif sum(comps) != monto:
                errores.append(f"fila {r} {medio}: el desglose {formula} suma {sum(comps)} y la celda dice {monto}")
            usd_v = sv.cell(r, col_usd).value
            usd = Decimal(repr(usd_v)) if isinstance(usd_v, (int, float)) else CERO
            items.append(ItemCaja(fecha, r, medio, banco, tipo, monto, usd, comps,
                                  formula if isinstance(formula, str) else "", modo, bancos))
    # Fechas repetidas o faltantes en el cierre (p. ej. dos filas «27/08» y ninguna «26/08»).
    filas_por_fecha: dict[dt.date, list[int]] = {}
    for r in range(f_sub + 1, sv.max_row + 1):
        f = sv.cell(r, c_fecha).value
        if isinstance(f, dt.datetime):
            filas_por_fecha.setdefault(correcciones.get(r, f.date()), []).append(r)
    for f, filas in sorted(filas_por_fecha.items()):
        if len(filas) > 1:
            errores.append(f"La fecha {f:%d/%m/%Y} aparece en {len(filas)} filas ({', '.join(map(str, filas))}).")
    if filas_por_fecha:
        d, fin = min(filas_por_fecha), max(filas_por_fecha)
        while d <= fin:
            if d not in filas_por_fecha:
                errores.append(f"No hay fila para el {d:%d/%m/%Y}.")
            d += dt.timedelta(days=1)
    faltantes = []
    if filas_por_fecha:
        d, fin = min(filas_por_fecha), max(filas_por_fecha)
        while d <= fin:
            if d not in filas_por_fecha:
                faltantes.append(d)
            d += dt.timedelta(days=1)
    return CierreCaja(ruta.name, hoja, tasas, items, errores, aplicadas, faltantes)

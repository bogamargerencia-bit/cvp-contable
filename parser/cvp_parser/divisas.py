"""Conciliación de las cuentas en divisas (efectivo $, Zelle, USDT, fondo de efectivo).

Tres fuentes, todas en US$:
  - Ventas: el cierre de caja diario trae lo recibido por día en cada medio («$ RECIBIDOS EN …»).
  - Libro del sistema por cuenta (Mayor Analítico): un asiento mensual de ventas («VTAS 08/2026»)
    y los demás movimientos (pagos, fondos, reparto, compra de dólares…).
  - Kardex de tesorería: ingresos del día por medio, pagos en divisas, otros movimientos, fondo de efectivo
    y saldos (VIENEN / DISPONIBLE).

Controles por cuenta:
  1. Ventas del mes = asiento de ventas del libro (o hasta un día, informando los que faltan).
  2. Ventas día a día = ingresos del día en el Kardex.
  3. Cada movimiento del libro (distinto de las ventas) contra el Kardex: exacto ±3 días; luego un asiento
     = suma de 2-3 movimientos del Kardex (o al revés); luego, si el Kardex lo metió en la fila de ventas
     del día, se explica con la diferencia de ese día.
  4. Saldo inicial y final del libro = VIENEN y DISPONIBLE del Kardex.
Nada se ajusta: lo que no cuadra queda como partida para la hoja Revisión.
"""
from __future__ import annotations

import datetime as dt
import itertools
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Optional

import openpyxl

from .kardex import Kardex, MovKardex
from .montos import CERO, monto_excel
from .sistema import AsientoLibro, LibroBanco


@dataclass
class CuentaDivisa:
    nombre: str                         # «Efectivo $», «Zelle», «USDT», «Fondo de efectivo»
    ventas_columna: Optional[str]       # regex del encabezado en el Excel de ventas (None = no recibe ventas)
    kardex_medio: Optional[str]         # DIVISAS / ZELLE / USDT / FONDO
    asiento_ventas: str = r"VTAS"       # regex (referencia o descripción) del asiento mensual de ventas


@dataclass
class DiaDivisa:
    fecha: dt.date
    ventas: Decimal
    kardex: Decimal
    nota: str = ""

    @property
    def diferencia(self) -> Decimal:
        return self.kardex - self.ventas


@dataclass
class PartidaDivisa:
    estado: str                       # Conciliado / Conciliado (agrupado) / Incluido en el día / Solo en libro / Solo en Kardex
    asientos: list[int] = field(default_factory=list)
    kardex: list[MovKardex] = field(default_factory=list)
    nota: str = ""


@dataclass
class ResultadoCuenta:
    cuenta: CuentaDivisa
    libro: LibroBanco
    ventas_mes: Optional[Decimal]
    asiento_ventas: Optional[int]
    ventas_nota: str
    ventas_ok: bool
    dias: list[DiaDivisa]
    partidas: list[PartidaDivisa]
    saldo_ini: tuple[Decimal, Optional[Decimal]]     # (libro, Kardex)
    saldo_fin: tuple[Decimal, Optional[Decimal]]

    @property
    def dias_con_diferencia(self) -> list[DiaDivisa]:
        return [d for d in self.dias if d.diferencia and "explicada" not in d.nota]


def leer_ventas_divisas(ruta: Path, cuentas: list[CuentaDivisa]) -> dict[str, dict[dt.date, Decimal]]:
    """Columnas «$ RECIBIDOS EN …» del Excel de ventas: {cuenta: {fecha: monto}}."""
    wb = openpyxl.load_workbook(str(ruta), data_only=True)
    ws = next(w for w in wb.worksheets if w.sheet_state == "visible" and w.title.upper().startswith("VENTAS"))
    out: dict[str, dict[dt.date, Decimal]] = {}
    for cta in cuentas:
        if not cta.ventas_columna:
            continue
        pos = next(((r, c.column) for r in range(1, 8) for c in ws[r]
                    if isinstance(c.value, str) and re.search(cta.ventas_columna, c.value, re.I)), None)
        if pos is None:
            raise ValueError(f"{Path(ruta).name}: no se encontró la columna de ventas «{cta.ventas_columna}»")
        f_enc, col = pos
        c_fecha = next(c.column for c in ws[f_enc] if str(c.value or "").strip() == "Fecha")
        dias: dict[dt.date, Decimal] = {}
        for r in range(f_enc + 1, ws.max_row + 1):
            f = ws.cell(r, c_fecha).value
            v = ws.cell(r, col).value
            if isinstance(f, dt.datetime) and isinstance(v, (int, float)):
                dias[f.date()] = dias.get(f.date(), CERO) + monto_excel(v)
        out[cta.nombre] = dias
    return out


def _signo(a: AsientoLibro) -> Decimal:
    return a.debito_usd - a.credito_usd


def conciliar_cuenta(cta: CuentaDivisa, libro: LibroBanco, ventas: Optional[dict[dt.date, Decimal]],
                     kardex: Optional[Kardex], dias_tol: int = 3) -> ResultadoCuenta:
    asientos = libro.asientos
    # 1. Ventas del mes contra el asiento de ventas.
    j_v = next((j for j, a in enumerate(asientos) if a.debito_usd and
                re.search(cta.asiento_ventas, f"{a.referencia} {a.descripcion}", re.I)), None)
    ventas_mes = sum(ventas.values(), CERO) if ventas is not None else None
    ventas_ok, nota_v = True, ""
    if ventas is not None:
        if j_v is None:
            ventas_ok, nota_v = False, f"El libro no tiene asiento de ventas; ventas del mes US$ {ventas_mes}."
        elif asientos[j_v].debito_usd == ventas_mes:
            nota_v = f"Asiento de ventas US$ {asientos[j_v].debito_usd} = ventas del mes."
        else:
            acum, hasta = CERO, None
            for f in sorted(ventas):
                acum += ventas[f]
                if acum == asientos[j_v].debito_usd:
                    hasta = f
            ventas_ok = False
            nota_v = (f"Asiento de ventas US$ {asientos[j_v].debito_usd} vs. ventas del mes US$ {ventas_mes} "
                      f"(diferencia {asientos[j_v].debito_usd - ventas_mes})"
                      + (f"; coincide con las ventas hasta el {hasta:%d/%m}." if hasta else "."))

    # 2. Ventas día a día contra el Kardex.
    dias: list[DiaDivisa] = []
    if ventas is not None and kardex is not None and cta.kardex_medio:
        kd = kardex.ventas_dia(cta.kardex_medio)
        for f in sorted(set(ventas) | set(kd)):
            dias.append(DiaDivisa(f, ventas.get(f, CERO), kd.get(f, CERO)))

    # 3. Movimientos del libro (sin el de ventas) contra el Kardex.
    partidas: list[PartidaDivisa] = []
    libres_a = [j for j in range(len(asientos)) if j != j_v]
    libres_k = list(kardex.no_ventas(cta.kardex_medio)) if kardex is not None and cta.kardex_medio else []

    def dias_(a: AsientoLibro, k: MovKardex) -> int:
        return abs((a.fecha - k.fecha).days)

    for j in sorted(libres_a, key=lambda j: asientos[j].fecha):            # exacto
        a = asientos[j]
        c = [k for k in libres_k if k.monto == _signo(a) and dias_(a, k) <= dias_tol]
        if c:
            k = min(c, key=lambda k: dias_(a, k))
            libres_k.remove(k)
            libres_a.remove(j)
            partidas.append(PartidaDivisa("Conciliado", [j], [k], f"Kardex fila {k.fila}: {k.descripcion}"))
    for j in list(libres_a):                                                # 1 asiento = 2-3 del Kardex
        a = asientos[j]
        c = [k for k in libres_k if dias_(a, k) <= dias_tol]
        grupo = next((g for n in (2, 3) for g in itertools.combinations(c, n)
                      if sum((k.monto for k in g), CERO) == _signo(a)), None)
        if grupo:
            for k in grupo:
                libres_k.remove(k)
            libres_a.remove(j)
            partidas.append(PartidaDivisa("Conciliado (agrupado)", [j], list(grupo),
                                          "1 asiento = " + " + ".join(f"{k.descripcion} {k.monto}" for k in grupo)))
    for k in list(libres_k):                                                # 1 del Kardex = 2-3 asientos
        c = [j for j in libres_a if dias_(asientos[j], k) <= dias_tol]
        grupo = next((g for n in (2, 3) for g in itertools.combinations(c, n)
                      if sum((_signo(asientos[j]) for j in g), CERO) == k.monto), None)
        if grupo:
            for j in grupo:
                libres_a.remove(j)
            libres_k.remove(k)
            partidas.append(PartidaDivisa("Conciliado (agrupado)", list(grupo), [k],
                                          f"{len(grupo)} asientos = Kardex fila {k.fila} {k.descripcion}"))
    por_dia = {d.fecha: d for d in dias}
    for j in list(libres_a):                                                # metido en la fila de ventas del día
        a = asientos[j]
        d = por_dia.get(a.fecha)
        if d and d.diferencia == _signo(a) and "explicada" not in d.nota:
            d.nota = f"Diferencia explicada: el Kardex incluye «{a.descripcion}» en la fila de ventas del día."
            libres_a.remove(j)
            partidas.append(PartidaDivisa("Incluido en el día del Kardex", [j], [],
                                          f"El Kardex suma {_signo(a)} en las ventas del {a.fecha:%d/%m}; "
                                          "no es una venta."))
    for j in libres_a:
        a = asientos[j]
        pista = next((k for k in libres_k if k.monto == _signo(a)), None)
        partidas.append(PartidaDivisa("Solo en libro", [j], [],
                                      f"Posible: Kardex fila {pista.fila} {pista.fecha:%d/%m} {pista.descripcion} "
                                      f"(fuera de ±{dias_tol} días)" if pista else ""))
    for k in libres_k:
        partidas.append(PartidaDivisa("Solo en Kardex", [], [k], ""))

    # 4. Saldos.
    ini_k = kardex.saldo_inicial.get(cta.kardex_medio) if kardex and cta.kardex_medio else None
    fin_k = kardex.disponible.get(cta.kardex_medio) if kardex and cta.kardex_medio else None
    return ResultadoCuenta(cta, libro, ventas_mes, j_v, nota_v, ventas_ok, dias, partidas,
                           (libro.saldo_inicial_usd, ini_k), (libro.saldo_final_usd, fin_k))


def conciliar_divisas(cuentas: list[CuentaDivisa], libros: dict[str, LibroBanco],
                      ventas: Optional[dict[str, dict[dt.date, Decimal]]],
                      kardex: Optional[Kardex]) -> list[ResultadoCuenta]:
    return [conciliar_cuenta(c, libros[c.nombre], (ventas or {}).get(c.nombre) if c.ventas_columna else None,
                             kardex) for c in cuentas if c.nombre in libros]

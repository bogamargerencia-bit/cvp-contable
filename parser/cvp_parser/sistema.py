"""Lector del libro de bancos exportado del sistema contable (SISTEMA_*.XLS).

Formato observado (WEI REST, agosto 2026; mismo esquema que el Mayor Analítico de OL BLUE):
  Fecha | Compro | Referencia | Cuenta | Debitos | Creditos | Saldo Mes | Saldo
  - Debitos/Creditos/Saldo están en US$.
  - "Referencia" trae el MONTO EN Bs. del movimiento bancario (p. ej. 61265 → 61.265,00).
    Cuando es texto ("PAGO MOVIL", "DEBITO", "VUELTOS") es un asiento resumen sin monto en Bs.
  - "Cuenta" trae en realidad la descripción del asiento.
  - Crédito en el libro = salida del banco (cargo). Débito en el libro = entrada (abono).
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Optional

import xlrd

from .montos import CERO, MontoInvalido, monto_excel, monto_us, monto_ve


@dataclass
class AsientoLibro:
    fila: int                         # fila en el Excel (1 = encabezado)
    fecha: dt.date
    comprobante: str
    referencia: str                   # tal como viene
    descripcion: str
    debito_usd: Decimal               # entrada al banco
    credito_usd: Decimal              # salida del banco
    saldo_usd: Optional[Decimal]
    monto_bs: Optional[Decimal]       # None si la referencia no es un monto
    referencia_banco: str = ""        # número de operación del banco, cuando la Referencia es eso
    # Montos en Bs. escritos en la descripción de un asiento resumen ("VUELTOS a + b + c").
    componentes_bs: list[Decimal] = field(default_factory=list)

    @property
    def es_salida(self) -> bool:
        return self.credito_usd > 0

    @property
    def monto_usd(self) -> Decimal:
        return self.credito_usd or self.debito_usd

    @property
    def tasa_implicita(self) -> Optional[Decimal]:
        if self.monto_bs is None or not self.monto_usd:
            return None
        return (self.monto_bs / self.monto_usd).quantize(Decimal("0.0001"))

    @property
    def es_resumen(self) -> bool:
        return self.monto_bs is None and not self.referencia_banco


@dataclass
class LibroBanco:
    archivo: str
    asientos: list[AsientoLibro]
    saldo_inicial_usd: Decimal
    saldo_final_usd: Decimal
    diferencias_saldo: list[str] = field(default_factory=list)
    filas_ignoradas: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)     # limitaciones del export (p. ej. sin saldos)


_COMPONENTE = re.compile(r"\d[\d.,]*\d")


def _componentes(texto: str) -> list[Decimal]:
    """'VUELTOS 22782.48 + 7184.20 + 970.30' -> [22782.48, 7184.20, 970.30]"""
    if "+" not in texto:
        return []
    out = []
    for t in _COMPONENTE.findall(texto):
        try:
            out.append(monto_excel(t))
        except MontoInvalido:
            return []
    return out if len(out) >= 2 else []


def _filas(ruta: Path) -> list[list]:
    """Filas del export como listas de 8 valores; la fecha (columna A) como datetime.date cuando lo es.
    Acepta el .xls original y un .xlsx (si el analista lo reguarda)."""
    if ruta.suffix.lower() == ".xlsx":
        import openpyxl
        ws = openpyxl.load_workbook(str(ruta), data_only=True).worksheets[0]
        filas = []
        for row in ws.iter_rows(values_only=True):
            v = ["" if c is None else c for c in (list(row) + [""] * 8)[:8]]
            if isinstance(v[0], dt.datetime):
                v[0] = v[0].date()
            filas.append(v)
        return filas
    wb = xlrd.open_workbook(str(ruta))
    sh = wb.sheet_by_index(0)
    filas = []
    for r in range(sh.nrows):
        v = (sh.row_values(r) + [""] * 8)[:8]
        if sh.ncols and sh.cell_type(r, 0) == xlrd.XL_CELL_DATE:
            v[0] = xlrd.xldate_as_datetime(v[0], wb.datemode).date()
        filas.append(v)
    return filas


ENCABEZADO = ["Fecha", "Compro", "Referencia", "Cuenta", "Debitos", "Creditos", "Saldo Mes", "Saldo"]


def _es_ref_bancaria(v: object) -> bool:
    """Referencia que es un número de operación del banco, no un monto: texto de 6+ dígitos
    ('083416407602') o un entero muy grande (≥ 1.000.000.000)."""
    if isinstance(v, str):
        return bool(re.fullmatch(r"\d{6,}", v.strip()))
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v).is_integer() and v >= 1_000_000_000
    return False


REPORTE = {"codcta", "nombre", "compro", "comentario", "fechatrans", "debitos", "creditos bs", "creditos $"}


def _celdas(ruta: Path) -> list[list]:
    """Todas las celdas de la primera hoja (fechas como datetime.date cuando la celda es fecha)."""
    if ruta.suffix.lower() == ".xlsx":
        import openpyxl
        ws = openpyxl.load_workbook(str(ruta), data_only=True).worksheets[0]
        return [[(c.date() if isinstance(c, dt.datetime) else ("" if c is None else c)) for c in row]
                for row in ws.iter_rows(values_only=True)]
    wb = xlrd.open_workbook(str(ruta))
    sh = wb.sheet_by_index(0)
    return [[xlrd.xldate_as_datetime(sh.cell_value(r, c), wb.datemode).date()
             if sh.cell_type(r, c) == xlrd.XL_CELL_DATE else sh.cell_value(r, c) for c in range(sh.ncols)]
            for r in range(sh.nrows)]


def _leer_reporte(ruta: Path, filas: list[list], f_enc: int) -> LibroBanco:
    """Export «reporte» (CACAO, septiembre 2026):
    codcta | nombre | compro | comentario | fechatrans | [fechacomp] | debitos | creditos BS | creditos $
    - Una fila inicial sin comprobante y en cero (apertura); no trae saldos.
    - creditos BS / creditos $ = salida del banco en Bs. y su equivalente en US$.
    - «debitos» no trae columna en US$: si viniera con montos no se puede interpretar y se informa.
    """
    col = {str(c).strip().lower(): i for i, c in enumerate(filas[f_enc])}
    asientos: list[AsientoLibro] = []
    ignoradas: list[str] = []
    for r in range(f_enc + 1, len(filas)):
        v = filas[r] + [""] * (len(col) - len(filas[r]))
        g = lambda k: v[col[k]]
        if not any(str(c).strip() for c in v):
            continue
        fecha = g("fechatrans")
        comp = str(g("compro")).strip()
        if not isinstance(fecha, dt.date):
            ignoradas.append(f"fila {r + 1}: sin fecha → {[c for c in v if c != '']}")
            continue
        deb = monto_excel(g("debitos") or 0)
        cre_bs = monto_excel(g("creditos bs") or 0)
        cre_usd = monto_excel(g("creditos $") or 0)
        if not comp and not deb and not cre_bs and not cre_usd:
            continue                                   # fila de apertura en cero
        if deb:
            raise ValueError(f"{ruta.name}, fila {r + 1}: el reporte trae un débito (entrada) de {deb} sin su "
                             "equivalente en US$; exporta el Mayor Analítico de la cuenta para procesarlo.")
        asientos.append(AsientoLibro(
            fila=r + 1, fecha=fecha, comprobante=comp, referencia=f"{cre_bs}", descripcion=str(g("comentario")).strip(),
            debito_usd=CERO, credito_usd=cre_usd, saldo_usd=None, monto_bs=cre_bs or None,
        ))
    total = sum((a.credito_usd for a in asientos), CERO)
    avisos = ["El libro viene en el formato «reporte», que no trae saldos: no se pueden verificar el saldo "
              "inicial ni el final del sistema."]
    if not any(a.debito_usd for a in asientos):
        avisos.append("El libro no trae ninguna entrada (débitos), solo salidas: todas las entradas del banco "
                      "quedarán «solo en banco». ¿Falta registrar los asientos de ventas o se exportó solo egresos?")
    return LibroBanco(archivo=ruta.name, asientos=asientos, saldo_inicial_usd=CERO, saldo_final_usd=-total,
                      filas_ignoradas=ignoradas, avisos=avisos)


_AMT_PDF = r"-?\d{1,3}(?:,\d{3})*\.\d{2}"
_FILA_PDF = re.compile(r"^(\d{2}/\d{2}/\d{4})\s+(\d+)\s+(\S+)\s+(.*?)\s*(" + _AMT_PDF + r")\s+(" + _AMT_PDF + r")\s*$")
_PIE_PDF = re.compile(r"Empresa:|Usuario:|Mayor Analitico|Desde la Fecha|Contacto:|Fecha:|Página:|^\s*J-\d")


def _leer_mayor_pdf(ruta: Path) -> LibroBanco:
    """Mayor Analítico impreso a PDF (Shiro: «MOVIMIENTOS DEL SISTEMA …» / «MAYOR DE MOVIMIENTOS …»).

    Mismas columnas que el export en Excel: Fecha | Compro | Referencia | Cuenta (descripción) | Debitos |
    Creditos | Saldo Mes | Saldo, en US$. Débito o crédito se decide por la POSICIÓN del monto bajo el
    encabezado de cada página; luego se verifica con el saldo de cada línea y con los «Totales:» impresos.
    La descripción puede seguir en las líneas siguientes (incluso en la página siguiente).
    Debe traer UNA sola cuenta: un mayor con varias cuentas (p. ej. todas las de gasto) se rechaza.
    """
    import subprocess
    t = subprocess.run(["pdftotext", "-layout", str(ruta), "-"], capture_output=True, text=True, check=True).stdout
    if "Mayor Analitico" not in t or not re.search(r"Debitos\s+Creditos", t):
        raise ValueError(f"{ruta.name}: es un PDF pero no es un Mayor Analítico del sistema; sube el export en "
                         "Excel o el Mayor Analítico de la cuenta en PDF.")
    cuentas = re.findall(r"Cuenta:\s+(\d+)\s+(.+?)\s{2,}Saldo Anterior:[ \t]*(" + _AMT_PDF + ")?", t)
    if len(cuentas) != 1:
        nombres = ", ".join(f"{c[0]} {c[1].strip()}" for c in cuentas[:4])
        raise ValueError(f"{ruta.name}: el Mayor trae {len(cuentas)} cuentas ({nombres}{'…' if len(cuentas) > 4 else ''}); "
                         "exporta el Mayor Analítico de UNA sola cuenta (la del banco o la divisa).")
    saldo_anterior = monto_us(cuentas[0][2]) if cuentas[0][2] else CERO

    asientos: list[AsientoLibro] = []
    ignoradas: list[str] = []
    difs: list[str] = []
    totales: list[Decimal] = []
    ult: Optional[AsientoLibro] = None
    n_linea = 0
    for pagina in t.split("\f"):
        cols: Optional[dict[str, int]] = None
        lineas = pagina.splitlines()
        i = 0
        while i < len(lineas):
            linea = lineas[i]
            n_linea += 1
            i += 1
            enc = re.search(r"Debitos\s+Creditos", linea)
            if linea.startswith("Fecha") and enc:
                cols = {"debe": linea.index("Debitos") + len("Debitos"), "haber": linea.index("Creditos") + len("Creditos")}
                continue
            if cols is None or not linea.strip() or linea.lstrip().startswith("Cuenta:") or _PIE_PDF.search(linea):
                continue
            if "Sub Total:" in linea or "Totales:" in linea:
                ult = None
                if "Totales:" in linea:
                    nums = re.findall(_AMT_PDF, linea)
                    while not nums and i < len(lineas):          # los montos van en la línea siguiente
                        nums = re.findall(_AMT_PDF, lineas[i]); i += 1; n_linea += 1
                    totales = [monto_us(x) for x in nums]
                continue
            f = _FILA_PDF.match(linea)
            if f and cols:
                monto, saldo = monto_us(f.group(5)), monto_us(f.group(6))
                fin = linea.index(f.group(5), f.start(5)) + len(f.group(5))
                es_debe = abs(fin - cols["debe"]) < abs(fin - cols["haber"])
                ref_raw = f.group(3)
                monto_bs, ref_banco = None, ""
                if re.fullmatch(r"\d+\.\d{2}", ref_raw):
                    monto_bs = Decimal(ref_raw)
                elif re.fullmatch(r"\d{1,3}(?:\.\d{3})*,\d{2}|\d+,\d{2}", ref_raw):
                    monto_bs = monto_ve(ref_raw)
                elif re.fullmatch(r"\d{6,}", ref_raw):
                    ref_banco = ref_raw
                ult = AsientoLibro(
                    fila=n_linea, fecha=dt.datetime.strptime(f.group(1), "%d/%m/%Y").date(), comprobante=f.group(2),
                    referencia=f"{monto_bs}" if monto_bs is not None else ref_raw,
                    descripcion=re.sub(r"\s+", " ", f.group(4)).strip(),
                    debito_usd=monto if es_debe else CERO, credito_usd=CERO if es_debe else monto,
                    saldo_usd=saldo, monto_bs=monto_bs, referencia_banco=ref_banco)
                asientos.append(ult)
            elif re.match(r"^\d{2}/\d{2}/\d{4}\s", linea):
                ignoradas.append(f"línea {n_linea}: no se pudo leer → {linea.strip()}")
            elif ult is not None and re.match(r"^\s{20,}\S", linea) and not re.search(_AMT_PDF, linea):
                ult.descripcion = f"{ult.descripcion} {linea.strip()}".strip(" .")
            else:
                ignoradas.append(f"línea {n_linea}: {linea.strip()}")
    for a in asientos:
        a.descripcion = re.sub(r"\s+\.$", "", re.sub(r"\s+", " ", a.descripcion)).strip()
        if a.monto_bs is None and not a.referencia_banco:
            a.componentes_bs = _componentes(a.descripcion)

    saldo = saldo_anterior
    for a in asientos:
        saldo = saldo + a.debito_usd - a.credito_usd
        if a.saldo_usd is not None and a.saldo_usd != saldo:
            difs.append(f"línea {a.fila} ({a.descripcion}): saldo calculado {saldo} vs. Mayor {a.saldo_usd}")
            saldo = a.saldo_usd
    deb = sum((a.debito_usd for a in asientos), CERO)
    cre = sum((a.credito_usd for a in asientos), CERO)
    if not totales:
        difs.append("No se encontraron los «Totales:» del Mayor; no se pudo verificar el total.")
    else:
        esperado = [x for x in (deb, cre) if x] + [saldo]
        faltan = [x for x in esperado if x not in totales and -x not in totales]
        if faltan:
            difs.append(f"Totales del Mayor {[str(x) for x in totales]} no coinciden con lo leído: débitos {deb}, "
                        f"créditos {cre}, saldo final {saldo}.")
    avisos = ["Mayor leído desde PDF (no del export en Excel)."]
    if not deb:
        avisos.append("El libro no registra ninguna entrada (débitos) en el período: todas las entradas del banco "
                      "quedarán «solo en banco». ¿Falta el asiento de ventas del cierre del mes?")
    return LibroBanco(archivo=ruta.name, asientos=asientos, saldo_inicial_usd=saldo_anterior, saldo_final_usd=saldo,
                      diferencias_saldo=difs, filas_ignoradas=ignoradas, avisos=avisos)


def leer_libro(ruta: Path) -> LibroBanco:
    """Lee el libro de bancos. Tres presentaciones del mismo sistema:
    - «reporte» (CACAO sept. 2026): codcta/nombre/compro/comentario/fechatrans/…/creditos BS/creditos $;
    - export simple (WEI REST): encabezado en la fila 1 y una fila por asiento;
    - Mayor Analítico completo (CACAO): bloque de empresa, línea «Cuenta: … Saldo Anterior: x»,
      encabezado repetido por página, descripciones que siguen en la línea siguiente,
      «Sub Total:» y «Totales:». Se usa el Saldo Anterior impreso y se verifican los Totales.
    """
    ruta = Path(ruta)
    with open(ruta, "rb") as fh:
        if fh.read(5) == b"%PDF-":
            return _leer_mayor_pdf(ruta)
    celdas = _celdas(ruta)
    f_rep = next((i for i, v in enumerate(celdas[:10]) if REPORTE <= {str(c).strip().lower() for c in v}), None)
    if f_rep is not None:
        return _leer_reporte(ruta, celdas, f_rep)
    filas = _filas(ruta)
    try:
        f_enc = next(i for i, v in enumerate(filas[:40]) if [str(c).strip() for c in v] == ENCABEZADO)
    except StopIteration:
        raise ValueError(f"{ruta.name}: no tiene un formato de libro conocido (Mayor Analítico, export simple "
                         f"o «reporte»). Encabezado esperado: {ENCABEZADO}")

    asientos: list[AsientoLibro] = []
    ignoradas: list[str] = []
    difs: list[str] = []
    saldo_anterior: Optional[Decimal] = None
    totales: Optional[tuple[Decimal, Decimal]] = None
    sin_debitos = False                        # Mayor sin columna de débitos (ver más abajo)
    avisos_corrimiento = False                 # se detectó por las filas, no por la línea «Cuenta:»
    totales_corridos: Optional[tuple[Decimal, Decimal]] = None
    ultimo_es_asiento = False
    for r in range(f_enc + 1, len(filas)):
        v = filas[r]
        textos = [str(c).strip() for c in v]
        if not any(textos):
            continue
        if not isinstance(v[0], dt.date):
            if textos == ENCABEZADO:                                  # encabezado de otra página
                continue
            if textos[0] == "Cuenta:":
                # El corrimiento se detecta por la POSICIÓN de la etiqueta, no por el valor: una cuenta que
                # arranca sin saldo en el sistema trae «Saldo Anterior:» vacío (Shiro, Bancamiga sept. 2026).
                if textos[6].startswith("Saldo Anterior"):
                    saldo_anterior = monto_excel(v[7]) if v[7] != "" else CERO
                elif textos[5].startswith("Saldo Anterior"):
                    # Mes sin débitos: el sistema omite esa columna y todo queda corrido a la izquierda.
                    saldo_anterior = monto_excel(v[6]) if v[6] != "" else CERO
                    sin_debitos = True
                continue
            if textos[3] in ("Sub Total:", "Totales:"):
                if textos[3] == "Totales:":
                    if sin_debitos:
                        totales_corridos = (monto_excel(v[4] or 0), monto_excel(v[6] or 0))
                    else:
                        totales = (monto_excel(v[4] or 0), monto_excel(v[5] or 0))
                continue
            if (asientos and ultimo_es_asiento and textos[3] and not any(textos[i] for i in (0, 1, 2, 4, 5, 6, 7))):
                asientos[-1].descripcion += " " + textos[3]         # la descripción sigue en esta línea
                continue
            if any(textos[1:]) and not re.match(r"(Empresa|Usuario|Desde la Fecha|Contacto|Fecha:)", " ".join(textos).strip()):
                ignoradas.append(f"fila {r + 1}: {[c for c in v if c != '']}")
            continue
        ultimo_es_asiento = True
        if not sin_debitos and v[7] == "" and v[6] != "" and v[5] == "" and v[4] != "":
            sin_debitos = True
            avisos_corrimiento = True
        fecha = v[0]
        comp = str(int(v[1])) if isinstance(v[1], (int, float)) else str(v[1]).strip()
        ref_raw = v[2]
        desc = str(v[3]).strip()
        monto_bs = None
        ref_banco = ""
        if _es_ref_bancaria(ref_raw):
            ref_banco = str(int(ref_raw)) if not isinstance(ref_raw, str) else ref_raw.strip()
            referencia = ref_banco
        elif isinstance(ref_raw, (int, float)) and not isinstance(ref_raw, bool) and ref_raw != 0:
            monto_bs = monto_excel(ref_raw)
            referencia = f"{monto_bs}"
        elif isinstance(ref_raw, str) and re.fullmatch(r"\d{1,3}(?:\.\d{3})*,\d{2}|\d+,\d{2}", ref_raw.strip()):
            monto_bs = monto_ve(ref_raw.strip())        # monto escrito como texto: «152637,69»
            referencia = f"{monto_bs}"
        elif isinstance(ref_raw, (int, float)) and not isinstance(ref_raw, bool):
            referencia = "0"          # Referencia 0: asiento resumen sin monto en Bs. (CACAO)
        else:
            referencia = str(ref_raw).strip()
        if sin_debitos:
            # Columnas corridas: [4] = monto, [6] = saldo. El sentido se decide por el saldo corrido.
            monto = monto_excel(v[4]) if v[4] != "" else CERO
            saldo_linea = monto_excel(v[6]) if v[6] != "" else None
            previo = asientos[-1].saldo_usd if asientos else saldo_anterior
            if previo is not None and saldo_linea == previo + monto:
                deb, cre = monto, CERO
            else:
                deb, cre = CERO, monto
                if previo is None or saldo_linea != previo - monto:
                    difs.append(f"fila {r + 1} ({desc}): no se pudo confirmar con el saldo si {monto} es entrada "
                                "o salida; se tomó como salida")
        else:
            deb = monto_excel(v[4]) if v[4] != "" else CERO
            cre = monto_excel(v[5]) if v[5] != "" else CERO
            saldo_linea = monto_excel(v[7]) if v[7] != "" else None
        asientos.append(AsientoLibro(
            fila=r + 1, fecha=fecha, comprobante=comp, referencia=referencia, descripcion=desc,
            debito_usd=deb,
            credito_usd=cre,
            saldo_usd=saldo_linea,
            monto_bs=monto_bs, referencia_banco=ref_banco,
            componentes_bs=_componentes(desc) if monto_bs is None and not ref_banco else [],
        ))

    # Saldo corrido en US$: desde el Saldo Anterior impreso o, si no hay, deducido de la primera línea.
    ini = CERO
    if saldo_anterior is not None:
        ini = saldo_anterior
    elif asientos and asientos[0].saldo_usd is not None:
        a = asientos[0]
        ini = a.saldo_usd - a.debito_usd + a.credito_usd
    saldo = ini
    for a in asientos:
        saldo = saldo + a.debito_usd - a.credito_usd
        if a.saldo_usd is not None and a.saldo_usd != saldo:
            difs.append(f"fila {a.fila} ({a.descripcion}): saldo calculado {saldo} vs. libro {a.saldo_usd}")
            saldo = a.saldo_usd
    if totales_corridos is not None:
        movido = sum((a.debito_usd + a.credito_usd for a in asientos), CERO)
        if (movido, saldo) != totales_corridos:
            difs.append(f"Totales del Mayor (sin columna de débitos): monto {totales_corridos[0]} / saldo "
                        f"{totales_corridos[1]} vs. suma de las líneas {movido} / saldo calculado {saldo}")
    avisos = []
    if avisos_corrimiento:
        avisos.append("Las columnas del Mayor vienen corridas (sin columna de débitos) y no se encontró la línea "
                      "«Cuenta: … Saldo Anterior»: se tomaron monto y saldo de las columnas corridas.")
    if sin_debitos:
        avisos.append("El Mayor no trae columna de débitos (el sistema la omite cuando el mes no tiene entradas): "
                      "el sentido de cada línea se tomó del saldo. El libro no registra ninguna entrada en el mes.")
    if totales is not None:
        deb = sum((a.debito_usd for a in asientos), CERO)
        cre = sum((a.credito_usd for a in asientos), CERO)
        if (deb, cre) != totales:
            difs.append(f"Totales del Mayor: débitos {totales[0]} / créditos {totales[1]} vs. suma de las líneas "
                        f"{deb} / {cre}")
    return LibroBanco(archivo=ruta.name, asientos=asientos, saldo_inicial_usd=ini,
                      saldo_final_usd=saldo, diferencias_saldo=difs, filas_ignoradas=ignoradas, avisos=avisos)

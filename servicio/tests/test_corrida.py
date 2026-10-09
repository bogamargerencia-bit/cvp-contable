"""Pruebas del servicio con un Supabase simulado en memoria y los archivos reales de agosto 2026
(si están en parser/tests/fixtures; si no, las pruebas con datos reales se saltan)."""
import hashlib
import uuid
from decimal import Decimal
from pathlib import Path

import pytest

from servicio.corrida import ErrorCorrida, armar_entrada, procesar_corrida

FIX = Path(__file__).resolve().parents[2] / "parser" / "tests" / "fixtures"
WEI = FIX / "wei_rest_2026_08"
CACAO = FIX / "cacao_2026_08"


class FakeSB:
    """Imita lo que usa procesar_corrida: tablas en memoria, storage en un dict."""

    def __init__(self):
        self.t = {k: [] for k in ("clientes", "periodos", "cuentas", "archivos", "corridas", "partidas",
                                  "decisiones", "correcciones_caja", "perfiles")}
        self.storage: dict[str, bytes] = {}

    @staticmethod
    def _ok(fila, filtros):
        for k, v in filtros.items():
            if k == "select":
                continue
            op, _, val = v.partition(".")
            if op == "eq" and str(fila.get(k)) != val:
                return False
            if op == "is" and fila.get(k) is not (val == "true"):
                return False
            if op == "in" and str(fila.get(k)) not in val.strip("()").split(","):
                return False
        return True

    def select(self, tabla, **f):
        filas = [dict(x) for x in self.t[tabla] if self._ok(x, f)]
        if tabla == "archivos":
            for a in filas:
                a["cuenta"] = next((c for c in self.t["cuentas"] if c["id"] == a["cuenta_id"]), None)
        return filas

    def uno(self, tabla, **f):
        r = self.select(tabla, **f)
        return r[0] if r else None

    def update(self, tabla, filtro, datos):
        for x in self.t[tabla]:
            if self._ok(x, filtro):
                x.update(datos)

    def reclamar(self, tabla, filtro, datos):
        n = [x for x in self.t[tabla] if self._ok(x, filtro)]
        for x in n:
            x.update(datos)
        return bool(n)

    def insert(self, tabla, filas):
        for f in filas:
            if tabla == "decisiones":            # valores por defecto de Postgres
                f = {"vigente": True, "decidido_en": "2026-10-08T00:00:00+00:00", **f}
            self.t[tabla].append(f)

    def descargar(self, ruta):
        return self.storage[ruta]

    def subir(self, ruta, contenido, tipo):
        self.storage[ruta] = contenido

    # ---- ayudas para armar casos
    def cliente(self, nombre, clave):
        c = {"id": str(uuid.uuid4()), "nombre": nombre, "clave_config": clave}
        p = {"id": str(uuid.uuid4()), "cliente_id": c["id"], "anio": 2026, "mes": 8, "estado": "abierto"}
        self.t["clientes"].append(c)
        self.t["periodos"].append(p)
        return c, p

    def cuenta(self, c, nombre, tipo="banco", banco=None):
        k = {"id": str(uuid.uuid4()), "cliente_id": c["id"], "nombre": nombre, "tipo": tipo, "banco": banco}
        self.t["cuentas"].append(k)
        return k

    def archivo(self, p, ruta: Path, tipo, cuenta=None, n=[0]):
        n[0] += 1
        datos = ruta.read_bytes()
        a = {"id": str(uuid.uuid4()), "periodo_id": p["id"], "cuenta_id": cuenta["id"] if cuenta else None,
             "tipo": tipo, "storage_path": f"x/{n[0]}-{ruta.name}", "nombre_original": ruta.name,
             "sha256": hashlib.sha256(datos).hexdigest(), "vigente": True, "subido_en": f"2026-10-06T10:{n[0]:02d}"}
        self.storage[a["storage_path"]] = datos
        self.t["archivos"].append(a)
        return a

    def corrida(self, p, numero=1):
        r = {"id": str(uuid.uuid4()), "periodo_id": p["id"], "numero": numero, "estado": "pendiente"}
        self.t["corridas"].append(r)
        return r


def test_falta_un_archivo(tmp_path):
    sb = FakeSB()
    c, p = sb.cliente("X", None)
    k = sb.cuenta(c, "BNC 1800", banco="BNC")
    sb.archivo(p, Path(__file__), "estado_cuenta", k)
    with pytest.raises(ErrorCorrida, match="falta el libro del sistema"):
        armar_entrada(c, p, sb.select("archivos"), sb.descargar, tmp_path)


def test_banco_sin_lector(tmp_path):
    sb = FakeSB()
    c, p = sb.cliente("X", None)
    k = sb.cuenta(c, "Activo", banco="ACTIVO")
    sb.archivo(p, Path(__file__), "estado_cuenta", k)
    with pytest.raises(ErrorCorrida, match="todavía no tiene lector"):
        armar_entrada(c, p, sb.select("archivos"), sb.descargar, tmp_path)


def test_archivo_alterado(tmp_path):
    sb = FakeSB()
    c, p = sb.cliente("X", None)
    k = sb.cuenta(c, "BNC", banco="BNC")
    a = sb.archivo(p, Path(__file__), "estado_cuenta", k)
    sb.archivo(p, Path(__file__), "libro_sistema", k)
    sb.storage[a["storage_path"]] += b"x"
    with pytest.raises(ErrorCorrida, match="SHA-256"):
        armar_entrada(c, p, sb.select("archivos"), sb.descargar, tmp_path)


def test_error_queda_en_la_corrida():
    sb = FakeSB()
    c, p = sb.cliente("X", None)
    r = sb.corrida(p)
    procesar_corrida(sb, r["id"])
    assert r["estado"] == "error" and "No hay ninguna cuenta bancaria" in r["error"]
    procesar_corrida(sb, r["id"])                      # ya no está pendiente: no hace nada
    assert r["estado"] == "error"


def test_periodo_cerrado():
    sb = FakeSB()
    c, p = sb.cliente("X", None)
    p["estado"] = "cerrado"
    r = sb.corrida(p)
    procesar_corrida(sb, r["id"])
    assert r["estado"] == "error" and "cerrado" in r["error"]


@pytest.mark.skipif(not (WEI / "kardex.xlsx").exists(), reason="faltan los archivos reales de WEI REST")
def test_wei_completo():
    sb = FakeSB()
    c, p = sb.cliente("ALIMENTOS SIERRA DEL SOL, C.A.", "WEI REST")
    for banco, edo, libro in [("100_BANCO", "edo_100_banco.pdf", "sistema_100_banco.xls"),
                              ("BNC", "edo_bnc.xls", "sistema_bnc.xls"),
                              ("BANPLUS", "edo_banplus.xlsx", "sistema_banplus.xls")]:
        k = sb.cuenta(c, banco, banco=banco)
        sb.archivo(p, WEI / edo, "estado_cuenta", k)
        sb.archivo(p, WEI / libro, "libro_sistema", k)
    for nombre, f in [("Efectivo $", "sistema_efectivo.xls"), ("Zelle", "sistema_zelle.xls"),
                      ("USDT", "sistema_usdt.xls"), ("Fondo de efectivo", "sistema_fondo_efectivo.xls")]:
        sb.archivo(p, WEI / f, "libro_sistema", sb.cuenta(c, nombre, tipo="divisa"))
    sb.archivo(p, WEI / "ventas.xlsx", "cierre_caja")
    sb.archivo(p, WEI / "kardex.xlsx", "kardex")
    r = sb.corrida(p)
    procesar_corrida(sb, r["id"])
    assert r["estado"] == "lista", r.get("error")
    assert (r["etapa"], r["avance"]) == ("Terminada", 100)
    res = r["resumen"]
    assert [b["banco"] for b in res["bancos"]] == ["100_BANCO", "BANPLUS", "BNC"]
    assert len(res["divisas"]) == 4 and all(d["ventas_ok"] for d in res["divisas"])
    assert res["uso"] == {"bancos": ["100_BANCO", "BANPLUS", "BNC"],
                          "divisas": ["Efectivo $", "Fondo de efectivo", "USDT", "Zelle"],
                          "cierre_caja": True, "kardex": True, "revision": False}
    assert r["ok_general"] is False and res["abiertas"] == len(sb.t["partidas"]) > 0
    assert len(r["archivos"]) == 12
    assert r["excel_path"].endswith("corridas/WEI_REST_2026-08_corrida01.xlsx")
    assert sb.storage[r["excel_path"]][:2] == b"PK"
    # Montos como texto exacto (nunca float).
    for x in sb.t["partidas"]:
        for m in ("monto_bs", "monto_usd"):
            assert x[m] is None or (isinstance(x[m], str) and Decimal(x[m]) == Decimal(x[m]).quantize(Decimal("0.01")))
    # Cada partida llega con su explicación y qué hacer (diagnóstico automático).
    assert all(x["explicacion"] and x["que_hacer"] for x in sb.t["partidas"])


@pytest.mark.skipif(not (CACAO / "ventas.xlsx").exists(), reason="faltan los archivos reales de CACAO")
def test_cacao_con_revision_anterior(tmp_path):
    """Corrida 1 → el Excel emitido se reintroduce como «revision» → corrida 2 lo reconoce."""
    sb = FakeSB()
    c, p = sb.cliente("CACAO COFFEE NG, C.A.", "CACAO")
    for banco, edo, libro in [("BANPLUS", "edo_banplus.xls", "sistema_banplus.xls"),
                              ("PLAZA", "edo_plaza.pdf", "sistema_plaza.xls")]:
        k = sb.cuenta(c, banco, banco=banco)
        sb.archivo(p, CACAO / edo, "estado_cuenta", k)
        sb.archivo(p, CACAO / libro, "libro_sistema", k)
    sb.archivo(p, CACAO / "ventas.xlsx", "cierre_caja")
    r1 = sb.corrida(p, 1)
    procesar_corrida(sb, r1["id"])
    assert r1["estado"] == "lista", r1.get("error")
    assert r1["resumen"]["corrida_revision"] == 1
    devuelto = tmp_path / "revision.xlsx"
    devuelto.write_bytes(sb.storage[r1["excel_path"]])
    sb.archivo(p, devuelto, "revision")
    r2 = sb.corrida(p, 2)
    procesar_corrida(sb, r2["id"])
    assert r2["estado"] == "lista", r2.get("error")
    assert r2["resumen"]["corrida_revision"] == 2 and r2["resumen"]["uso"]["revision"] is True
    # Sin decisiones en el Excel: todo sigue abierto como «Sin decisión».
    assert set(r2["resumen"]["situaciones"]) == {"Sin decisión"}


def _decidir(sb, p, corrida, codigo, decision, comentario="", quien="u-analista", cuando="2026-10-07T12:00:00+00:00"):
    """Lo que hace la función registrar_decision de la web: la anterior deja de ser vigente."""
    for d in sb.t["decisiones"]:
        if d["codigo"] == codigo and d["vigente"]:
            d["vigente"] = False
    sb.t["decisiones"].append({"periodo_id": p["id"], "codigo": codigo, "decision": decision,
                               "comentario": comentario or None, "revisado_por": quien, "decidido_en": cuando,
                               "corrida_id": corrida["id"], "vigente": True})


def _cacao(sb):
    c, p = sb.cliente("CACAO COFFEE NG, C.A.", "CACAO")
    for banco, edo, libro in [("BANPLUS", "edo_banplus.xls", "sistema_banplus.xls"),
                              ("PLAZA", "edo_plaza.pdf", "sistema_plaza.xls")]:
        k = sb.cuenta(c, banco, banco=banco)
        sb.archivo(p, CACAO / edo, "estado_cuenta", k)
        sb.archivo(p, CACAO / libro, "libro_sistema", k)
    sb.archivo(p, CACAO / "ventas.xlsx", "cierre_caja")
    sb.t["perfiles"].append({"id": "u-analista", "nombre": "Margareth Celis", "email": "m@x.test"})
    return c, p


@pytest.mark.skipif(not (CACAO / "ventas.xlsx").exists(), reason="faltan los archivos reales de CACAO")
def test_decisiones_en_linea():
    """Corrida 1 → decisiones en la web (sin Excel) → corrida 2 las aplica; la corrección de caja aceptada
    se guarda y se sigue aplicando en la corrida 3."""
    sb = FakeSB()
    c, p = _cacao(sb)
    r1 = sb.corrida(p, 1)
    procesar_corrida(sb, r1["id"])
    assert r1["estado"] == "lista", r1.get("error")
    ps = [x for x in sb.t["partidas"] if x["corrida_id"] == r1["id"]]
    prop = next(x for x in ps if x["propuesta"])
    otras = [x for x in ps if not x["propuesta"]][:3]
    _decidir(sb, p, r1, prop["codigo"], "Aceptar")
    _decidir(sb, p, r1, otras[0]["codigo"], "Justificado", "Verificado con el banco")
    _decidir(sb, p, r1, otras[1]["codigo"], "Justificado")          # sin comentario: queda incompleta
    r2 = sb.corrida(p, 2)
    r2["creada_por"] = "u-analista"
    procesar_corrida(sb, r2["id"])
    assert r2["estado"] == "lista", r2.get("error")
    p2 = {x["codigo"]: x for x in sb.t["partidas"] if x["corrida_id"] == r2["id"]}
    assert prop["codigo"] not in p2                       # la fecha se corrigió: la propuesta desaparece
    assert p2[otras[0]["codigo"]]["situacion"] == "Cerrado"
    assert p2[otras[1]["codigo"]]["situacion"] == "Decisión incompleta"
    assert "comentario" in p2[otras[1]["codigo"]]["aviso"]
    assert p2[otras[2]["codigo"]]["situacion"] == "Sin decisión"
    [corr] = sb.t["correcciones_caja"]
    assert corr["fila"] == prop["propuesta"]["fila"] and corr["fecha_nueva"] == prop["propuesta"]["fecha"]
    assert corr["codigo"] == prop["codigo"] and corr["aceptada_por"] == "u-analista"
    hist = r2["resumen"]["historial"]
    assert any(h[2] == otras[0]["codigo"] and h[7] == "Margareth Celis" for h in hist)
    # Corrida 3: la corrección sigue aplicada (viene de la base) y el historial se conserva.
    r3 = sb.corrida(p, 3)
    procesar_corrida(sb, r3["id"])
    assert r3["estado"] == "lista", r3.get("error")
    p3 = {x["codigo"]: x for x in sb.t["partidas"] if x["corrida_id"] == r3["id"]}
    assert prop["codigo"] not in p3 and p3[otras[0]["codigo"]]["situacion"] == "Cerrado"
    assert len(sb.t["correcciones_caja"]) == 1
    assert len(r3["resumen"]["historial"]) >= len(hist)


@pytest.mark.skipif(not (CACAO / "ventas.xlsx").exists(), reason="faltan los archivos reales de CACAO")
def test_excel_y_web_gana_lo_mas_reciente(tmp_path):
    import openpyxl
    sb = FakeSB()
    c, p = _cacao(sb)
    r1 = sb.corrida(p, 1)
    procesar_corrida(sb, r1["id"])
    ps = [x for x in sb.t["partidas"] if x["corrida_id"] == r1["id"] and not x["propuesta"]]
    a, b = ps[0]["codigo"], ps[1]["codigo"]
    # En la web: «a» antes de subir el Excel, «b» después.
    _decidir(sb, p, r1, a, "Justificado", "web antes", cuando="2026-10-06T09:00:00+00:00")
    _decidir(sb, p, r1, b, "Justificado", "web después", cuando="2026-10-07T23:00:00+00:00")
    wb = openpyxl.load_workbook(__import__("io").BytesIO(sb.storage[r1["excel_path"]]))
    ws = wb["Revisión"]
    fila_enc = next(r for r in range(1, 40) if ws.cell(r, 1).value == "Código")
    enc = {ws.cell(fila_enc, k).value: k for k in range(1, ws.max_column + 1)}
    for r in range(fila_enc + 1, ws.max_row + 1):
        if ws.cell(r, 1).value in (a, b):
            ws.cell(r, enc["Decisión"], "Justificado")
            ws.cell(r, enc["Comentario"], "excel")
            ws.cell(r, enc["Revisado por"], "MC")
    ruta = tmp_path / "rev.xlsx"
    wb.save(ruta)
    arch = sb.archivo(p, ruta, "revision")
    arch["subido_en"], arch["subido_por"] = "2026-10-07T10:00:00+00:00", "u-analista"
    r2 = sb.corrida(p, 2)
    procesar_corrida(sb, r2["id"])
    assert r2["estado"] == "lista", r2.get("error")
    vig = {d["codigo"]: d for d in sb.t["decisiones"] if d["vigente"]}
    assert vig[a]["comentario"] == "excel"            # el Excel es más reciente que la decisión en la web
    assert vig[b]["comentario"] == "web después"      # la web es más reciente que el Excel
    assert len([d for d in sb.t["decisiones"] if d["codigo"] == a]) == 2


def test_cuenta_inactiva_se_ignora(tmp_path):
    sb = FakeSB()
    c, p = sb.cliente("X", None)
    k = sb.cuenta(c, "BNC equivocada", banco="BNC")
    k["activo"] = False
    sb.archivo(p, Path(__file__), "estado_cuenta", k)
    archivos = sb.select("archivos", periodo_id=f"eq.{p['id']}")
    with pytest.raises(ErrorCorrida, match="No hay ninguna cuenta bancaria"):
        armar_entrada(c, p, archivos, sb.descargar, tmp_path)


@pytest.mark.skipif(not (WEI / "ventas.xlsx").exists(), reason="faltan los archivos reales de WEI")
def test_solo_conversion():
    """Cliente en modo conversión: solo estados de cuenta, sin libro; Excel con resumen por ítem."""
    import io
    import openpyxl
    sb = FakeSB()
    c, p = sb.cliente("CLIENTE CONVERSION", None)
    c["config"] = {"modo": "conversion"}
    for banco, edo in [("BNC", "edo_bnc.xls"), ("BANPLUS", "edo_banplus.xlsx"), ("100_BANCO", "edo_100_banco.pdf")]:
        sb.archivo(p, WEI / edo, "estado_cuenta", sb.cuenta(c, f"{banco} principal", banco=banco))
    sb.cuenta(c, "BNC sin archivo", banco="BNC")          # cuenta sin estado de cuenta: no impide procesar
    r = sb.corrida(p, 1)
    procesar_corrida(sb, r["id"])
    assert r["estado"] == "lista", r.get("error")
    res = r["resumen"]
    assert res["modo"] == "conversion" and res["ok_general"] is True and len(res["cuentas"]) == 3
    assert not sb.t["partidas"]
    bnc = next(x for x in res["cuentas"] if x["banco"] == "BNC")
    assert bnc["movimientos"] == 794 and bnc["debitos"] == "50363951.11" and bnc["creditos"] == "50379929.63"
    wb = openpyxl.load_workbook(io.BytesIO(sb.storage[r["excel_path"]]))
    assert wb.sheetnames[:2] == ["Resumen", "Resumen por ítem"]


def test_solo_conversion_sin_estados():
    sb = FakeSB()
    c, p = sb.cliente("X", None)
    c["config"] = {"modo": "conversion"}
    sb.cuenta(c, "BNC", banco="BNC")
    r = sb.corrida(p, 1)
    procesar_corrida(sb, r["id"])
    assert r["estado"] == "error" and "ningún estado de cuenta" in r["error"]
